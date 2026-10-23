# PROPRIETARY AND CONFIDENTIAL
#
# This file and its contents are the proprietary property of LGERC HW ATSC 60Hz.
# Unauthorized copying, redistribution, modification, disclosure, or claim of
# ownership is strictly prohibited.
#
# Use of LLMs, AI agents, AI chatbots, or automated code tools to alter,
# redefine, disguise, or claim this source code as another party’s own work is
# prohibited and may result in legal action.
#
# Removal or deletion of this notice, whether intentional or unintentional,
# does not waive any rights and shall be considered a breach of the applicable
# repository license notice.
#
# ALL RIGHTS RESERVED.



#from streamlit_navigation_bar import st_navbar
import pandas as pd
import time
import requests
import json
import io
import numpy as np
from openpyxl import load_workbook, workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import polars as pl
from pathlib import Path
import os
from dej_sdk import bucket
import tempfile
import re
import polars as pl
import hashlib
from pandas.io.formats.style import Styler
import connectorx

from main_function.services.check_sso import SSO_NPDM
from main_function.services.main_hwpart import targetpart_compare, developmentpart_compare, initialpart_compare
from main_function.services.main_tooloption import run_tooloption, getToolinfo
from main_function.services.subs_hwpart import run_subshwpart
from datetime import timedelta, datetime, date

from main_function.functions.get_latestver import read_latest_version
from main_function.functions.getdata import HerosURL, userURL, CoworkURL, OPurl

def get_suffix_data(suffixCode: str):

        login = SSO_NPDM()
        login.auto_login()

        login_bearer = login.auth_bearer


        suffixUrl = "https://npdm.lge.com/api/modelDev/search-suffix-rule"
        RefererHost = "https://npdm.lge.com/automation/bom/suffix-master"


        headers = {
            "Accept": "application/json, text/plain, */*",
            "Authorization": login_bearer,
            "BusinessUnit": "HE",
            "Content-Type": "application/json",
            "Host": "npdm.lge.com",
            "Referer": RefererHost
        }

        params = {
            "searchVO": {
                "affiliateCode": "*",
                "divisionCode": "GLZ",
                "plantUnitCode": "GLZ.EKHQ",
                "productCode": "PDT002",
                "year": "",
                "divOption": suffixCode,
                "states": "Active"
            }
        }

        response = requests.put(
            suffixUrl,
            headers=headers,
            json=params,
            timeout=60
        )

        get_status = response.status_code

        print("suffixCode:", suffixCode)
        print("status:", get_status)
        print("response preview:", response.text[:1000])

        if get_status != 200:
            return get_status, None, pl.DataFrame()

        get_json = response.json()
        #get_data_panel = get_json.get("data", [])
        get_data_panel = get_json["data"]

        if not get_data_panel:
            return get_status, get_json, pl.DataFrame()

        panelData = pl.DataFrame(get_data_panel, infer_schema_length=1000)

        return get_status, get_json, panelData


