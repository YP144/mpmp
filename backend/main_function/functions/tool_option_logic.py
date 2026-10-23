# PROPRIETARY AND CONFIDENTIAL
#
# This file and its contents are the proprietary property of LGERC HW ATSC 60Hz.
# Unauthorized copying, redistribution, modification, disclosure, or claim of
# ownership is strictly prohibited.
#
# Use of LLMs, AI agents, AI chatbots, or automated code tools to alter,
# redefine, disguise, or claim this source code as another party’s own work is
# prohibited and may result in legal action.
#
# Removal or deletion of this notice, whether intentional or unintentional,
# does not waive any rights and shall be considered a breach of the applicable
# repository license notice.
#
# ALL RIGHTS RESERVED.



#from streamlit_navigation_bar import st_navbar
import pandas as pd
import time
import requests
import json
import io
import numpy as np
from openpyxl import load_workbook, workbook
from openpyxl import Workbook 
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import polars as pl
from pathlib import Path
import os
from dej_sdk import bucket
import tempfile
import re
import polars as pl
import hashlib
from pandas.io.formats.style import Styler
import connectorx
from copy import copy as copy_style
import zipfile
import re
import threading

from main_function.functions.support import set_cell

from main_function.services.check_sso import SSO_NPDM
from main_function.services.main_tooloption import (
    target_tooloption, development_tooloption, initial_tooloption,
    calculate_tooloption1, calculate_tooloption2, calculate_tooloption3,
    calculate_tooloption4, calculate_tooloption5, calculate_tooloption6,
    run_tooloption
    )


def compare_tool_option_summary(BOMframe, target_model, bearerAuth):

    try:
        ebt_tool_option = ToolOptionCross(BOMframe).rename(
            {"settings": "mp_settings"}
        )
        master_tool_option = GetModelToolOptionMaster(
            target_model,
            bearerAuth
        )

        ebt_tool_option = pl.DataFrame(ebt_tool_option)
        master_tool_option = pl.DataFrame(master_tool_option)

        result_tool_option = pl.DataFrame({
            "No": list(range(1, 12)),
            "Tool Option": [
                "Board Option",
                "Tool Option 1",
                "Tool Option 2",
                "Tool Option 3",
                "Tool Option 4",
                "Tool Option 5",
                "Tool Option 6",
                "Tool CRC",
                "Country Group",
                "Country",
                "Area Code"
            ],
            "MP Tool Option": [None] * 11,
            "Master Tool Option": [None] * 11,
            "Remark": [None] * 11,
        })

        # MP Tool Option
        result_tool_option = set_cell(result_tool_option, 0, "MP Tool Option", ebt_tool_option.item(8, "mp_settings"))
        result_tool_option = set_cell(result_tool_option, 1, "MP Tool Option", ebt_tool_option.item(1, "mp_settings"))
        result_tool_option = set_cell(result_tool_option, 2, "MP Tool Option", ebt_tool_option.item(2, "mp_settings"))
        result_tool_option = set_cell(result_tool_option, 3, "MP Tool Option", ebt_tool_option.item(3, "mp_settings"))
        result_tool_option = set_cell(result_tool_option, 4, "MP Tool Option", ebt_tool_option.item(4, "mp_settings"))
        result_tool_option = set_cell(result_tool_option, 5, "MP Tool Option", ebt_tool_option.item(5, "mp_settings"))
        result_tool_option = set_cell(result_tool_option, 6, "MP Tool Option", ebt_tool_option.item(6, "mp_settings"))

        result_tool_option = set_cell(result_tool_option, 7, "MP Tool Option", ebt_tool_option.item(9, "mp_settings"))
        result_tool_option = set_cell(result_tool_option, 8, "MP Tool Option", ebt_tool_option.item(11, "mp_settings"))
        result_tool_option = set_cell(result_tool_option, 9, "MP Tool Option", ebt_tool_option.item(12, "mp_settings"))
        result_tool_option = set_cell(result_tool_option, 10, "MP Tool Option", ebt_tool_option.item(10, "mp_settings"))

        # Master Tool Option
        result_tool_option = set_cell(result_tool_option, 0, "Master Tool Option", master_tool_option.item(17, "settings"))
        result_tool_option = set_cell(result_tool_option, 1, "Master Tool Option", master_tool_option.item(5, "settings"))
        result_tool_option = set_cell(result_tool_option, 2, "Master Tool Option", master_tool_option.item(6, "settings"))
        result_tool_option = set_cell(result_tool_option, 3, "Master Tool Option", master_tool_option.item(7, "settings"))
        result_tool_option = set_cell(result_tool_option, 4, "Master Tool Option", master_tool_option.item(8, "settings"))
        result_tool_option = set_cell(result_tool_option, 5, "Master Tool Option", master_tool_option.item(9, "settings"))
        result_tool_option = set_cell(result_tool_option, 6, "Master Tool Option", master_tool_option.item(10, "settings"))

        result_tool_option = set_cell(result_tool_option, 7, "Master Tool Option", master_tool_option.item(13, "settings"))
        result_tool_option = set_cell(result_tool_option, 8, "Master Tool Option", master_tool_option.item(15, "settings"))
        result_tool_option = set_cell(result_tool_option, 9, "Master Tool Option", master_tool_option.item(14, "settings"))
        result_tool_option = set_cell(result_tool_option, 10, "Master Tool Option", master_tool_option.item(16, "settings"))

        result_tool_option = result_tool_option.with_columns(
            pl.when(
                pl.col("MP Tool Option") == pl.col("Master Tool Option")
            )
            .then(pl.lit("OK"))
            .otherwise(pl.lit("NG"))
            .alias("Status")
        )

        result_tool_option = result_tool_option.with_columns(
            pl.lit("").alias("Remarks")
        )

        result_check = (
            "NG"
            if result_tool_option["Status"].eq("NG").any()
            else "OK"
        )

        ebt_tool_json = result_tool_option.to_dicts()

        return (
            result_tool_option,
            master_tool_option,
            ebt_tool_json,
            result_check
        )

    except Exception as e:

        result_tool_option = pl.DataFrame({
            "No": list(range(1, 12)),
            "Tool Option": [
                "Board Option",
                "Tool Option 1",
                "Tool Option 2",
                "Tool Option 3",
                "Tool Option 4",
                "Tool Option 5",
                "Tool Option 6",
                "Tool CRC",
                "Country Group",
                "Country",
                "Area Code"
            ],
            "MP Tool Option": ["No Tool Option Data"] * 11,
            "Master Tool Option": ["No Tool Option Data"] * 11,
            "Status": ["CHECK"] * 11,
            "Remarks": [f"Error: {str(e)}"] * 11
        })

        ebt_tool_json = result_tool_option.to_dicts()

        return (
            result_tool_option,
            pl.DataFrame(),
            ebt_tool_json,
            "CHECK"
        )

def compar_tool_option_list(result_tool_option):
    #mp bom tool option list
    mp_tool1 = result_tool_option.item[1, "MP Tool Option"]
    mp_tool2 = result_tool_option.item[2, "MP Tool Option"]
    mp_tool3 = result_tool_option.item[3, "MP Tool Option"]
    mp_tool4 = result_tool_option.item[4, "MP Tool Option"]
    mp_tool5 = result_tool_option.item[5, "MP Tool Option"]
    mp_tool6 = result_tool_option.item[6, "MP Tool Option"]
    mp_toolCRC = result_tool_option.item[7, "MP Tool Option"]
    mp_country = result_tool_option.item[8, "MP Tool Option"]
    mp_areacode = result_tool_option.item[9, "MP Tool Option"]

    mp_boardOpt = result_tool_option.item[1, "MP Tool Option"]

    #master bom option list





def ToolOptionCross(BOMframe):

    default_message = "Job explanation not exist"

    def make_result_frame(result_dict):
        return pl.DataFrame({
            "tool_option": list(result_dict.keys()),
            "settings": [
                value.write_json() if isinstance(value, pl.DataFrame)
                else default_message if value is None
                else str(value)
                for value in result_dict.values()
            ]
        })

    def default_result(getEBTrow=None):
        if isinstance(getEBTrow, pl.DataFrame):
            getEBTrow_value = getEBTrow.write_json()
        else:
            getEBTrow_value = default_message

        return make_result_frame({
            "ToolGet": default_message,
            "Tool1": default_message,
            "Tool2": default_message,
            "Tool3": default_message,
            "Tool4": default_message,
            "Tool5": default_message,
            "Tool6": default_message,
            "Tool7": default_message,
            "BoardOpt": default_message,
            "CRCTool": default_message,
            "AreaCode": default_message,
            "CountryGroup": default_message,
            "GetToolCountry": default_message,
            "getEBTrow": getEBTrow_value
        })

    getEBTrow = BOMframe.filter(
        (pl.col("part_number").cast(pl.Utf8).str.contains(r"^EBT")) &
        (pl.col("bomLevel") == 1)
    )

    if getEBTrow.height == 0:
        return default_result()

    if "JobExpl" not in getEBTrow.columns:
        return default_result(getEBTrow)

    SentenceTool = getEBTrow.item(0, "JobExpl")

    if SentenceTool is None or str(SentenceTool).strip() == "":
        return default_result(getEBTrow)

    SentenceTool = str(SentenceTool)

    try:
        ToolOptionModel = {}

        BoardOpt_match = re.search(r"Board Option\s*:\s*([^/]+)", SentenceTool)
        ToolOpt_match = re.search(r"Tool Option\s*\((.*?)\)", SentenceTool)
        CRC_match = re.search(r"Board\+ToolCRC\s*:\s*([^/]+)", SentenceTool)
        AreaCode_match = re.search(r"Area Code\s*:\s*([^/]+)", SentenceTool)
        CountryGroup_match = re.search(r"Country Group\s*:\s*([^/]+)", SentenceTool)
        Country_match = re.search(r"Country\s*:\s*([^/]+)", SentenceTool)

        if BoardOpt_match:
            ToolOptionModel["Board Option"] = BoardOpt_match.group(1).strip()

        if ToolOpt_match:
            ToolOptionModel["Tool Option"] = ToolOpt_match.group(1).strip()

        if CRC_match:
            ToolOptionModel["Board+ToolCRC"] = CRC_match.group(1).strip()

        if AreaCode_match:
            ToolOptionModel["Area Code"] = AreaCode_match.group(1).strip()

        if CountryGroup_match:
            ToolOptionModel["Country Group"] = CountryGroup_match.group(1).strip()

        if Country_match:
            ToolOptionModel["Country"] = Country_match.group(1).strip()

        if "Tool Option" not in ToolOptionModel:
            return default_result(getEBTrow)

        ToolGet = (
            pl.DataFrame({
                "tool_option_raw": [ToolOptionModel["Tool Option"]]
            })
            .with_columns(
                pl.col("tool_option_raw").str.split(",").alias("tool_option")
            )
            .explode("tool_option")
            .with_columns(
                pl.col("tool_option").str.split(":").alias("split_pair")
            )
            .with_columns(
                pl.col("split_pair").list.get(0).cast(pl.Int64).alias("tool_number"),
                pl.col("split_pair").list.get(1).cast(pl.Int64).alias("tool_value")
            )
            .select(["tool_number", "tool_value"])
        )

        tool_map = dict(zip(
            ToolGet["tool_number"].to_list(),
            ToolGet["tool_value"].to_list()
        ))

        Tool1 = tool_map.get(1)
        Tool2 = tool_map.get(2)
        Tool3 = tool_map.get(3)
        Tool4 = tool_map.get(4)
        Tool5 = tool_map.get(5)
        Tool6 = tool_map.get(6)
        Tool7 = tool_map.get(7)

        return make_result_frame({
            "ToolGet": ToolGet,
            "Tool1": Tool1,
            "Tool2": Tool2,
            "Tool3": Tool3,
            "Tool4": Tool4,
            "Tool5": Tool5,
            "Tool6": Tool6,
            "Tool7": Tool7,
            "BoardOpt": ToolOptionModel.get("Board Option", default_message),
            "CRCTool": ToolOptionModel.get("Board+ToolCRC", default_message),
            "AreaCode": ToolOptionModel.get("Area Code", default_message),
            "CountryGroup": ToolOptionModel.get("Country Group", default_message),
            "GetToolCountry": ToolOptionModel.get("Country", default_message),
            "getEBTrow": getEBTrow
        })

    except Exception as e:
        return make_result_frame({
            "ToolGet": f"Parse error: {str(e)}",
            "Tool1": default_message,
            "Tool2": default_message,
            "Tool3": default_message,
            "Tool4": default_message,
            "Tool5": default_message,
            "Tool6": default_message,
            "Tool7": default_message,
            "BoardOpt": default_message,
            "CRCTool": default_message,
            "AreaCode": default_message,
            "CountryGroup": default_message,
            "GetToolCountry": default_message,
            "getEBTrow": getEBTrow
        })


def GetModelToolOptionMaster(targetModel, bearerAuth):

    #sso = SSO_NPDM()
    #is_logged_in = sso.auto_login(ssolgenet_cookie, smsession_cookie)

    #bearerAuth = sso.auth_bearer

    NO_DATA = "No data Detected"

    def value_to_string(value):
        if isinstance(value, pl.DataFrame):
            return value.write_json()

        if isinstance(value, dict) or isinstance(value, list):
            return json.dumps(value, ensure_ascii=False, default=str)

        if value is None:
            return NO_DATA

        return str(value)

    def master_tool_option(result_dict):
        return pl.DataFrame({
            "tool_option": list(result_dict.keys()),
            "settings": [
                value_to_string(value)
                for value in result_dict.values()
            ]
        })

    """def no_data_result(DataToolOption=None, ToolOptionFrame=None, GetToolOption=None):
        return master_tool_option({
            "DataToolOption": DataToolOption,
            "ToolOptionFrame": ToolOptionFrame,
            "GetToolOption": GetToolOption,
            "ModelName": NO_DATA,
            "ModelSuffix": NO_DATA,
            "AppliedYear": NO_DATA,
            "tool1": NO_DATA,
            "tool2": NO_DATA,
            "tool3": NO_DATA,
            "tool4": NO_DATA,
            "tool5": NO_DATA,
            "tool6": NO_DATA,
            "tool7": NO_DATA,
            "commercialTool": NO_DATA,
            "crcSum": NO_DATA,
            "country": NO_DATA,
            "countryGroup": NO_DATA,
            "areaCode": NO_DATA,
            "boardOption": NO_DATA
        })"""

    targetModel = targetModel.split(".")
    GetModelName = targetModel[0]

    headers = {
        "Accept": "application/json, text/plain, */*",
        "Authorization": bearerAuth,
        "BusinessUnit": "HE",
        "Content-Type": "application/json",
        "Host": "npdm.lge.com",
        "Referer": "https://npdm.lge.com/my-desk/tool-option"
    }

    GetToolOptionParams = {
        "cParamToolOptionVO": {
            "modelName": f"{targetModel}"
        },
        "pagingEntity": {
            "omfOrderByBuilder": {
                "isEnd": "false",
                "isOQL": "true",
                "list": [],
                "listType": []
            },
            "totalCount": 1000,
            "targetRow": 1,
            "rowSize": 1000,
            "currentPage": 1,
            "executeCount": "true"
        }
    }

    response = requests.put(
        "https://npdm.lge.com/api/change/tool/option/list",
        headers=headers,
        json=GetToolOptionParams
    )

    DataToolOption = response.json()

    EarnToolOpt = DataToolOption.get("data") or {}
    all_outdata = EarnToolOpt.get("list") or []

    if not all_outdata:
        ToolOptionFrame = pl.DataFrame()
        GetToolOption = pl.DataFrame()

        return no_data_result(
            DataToolOption=DataToolOption,
            ToolOptionFrame=ToolOptionFrame,
            GetToolOption=GetToolOption
        )

    ToolOptionFrame = pl.DataFrame(all_outdata, infer_schema_length=None)

    if ToolOptionFrame.is_empty():
        GetToolOption = pl.DataFrame()

        return no_data_result(
            DataToolOption=DataToolOption,
            ToolOptionFrame=ToolOptionFrame,
            GetToolOption=GetToolOption
        )

    ToolOptionFrame = ToolOptionFrame.with_columns(
        pl.col("modelName").cast(pl.Utf8).str.replace_all(r"_", "."),
        pl.col("modelSuffix").cast(pl.Utf8).str.replace_all(r"_", ".")
    )
    #matchesSuffix = matchesSuffix.row(0)

    """GetToolOption = ToolOptionFrame.filter(
        pl.col("modelSuffix").str.contains_or_any(GetSuffixName)
    )"""

    """if GetToolOption.is_empty():
        return no_data_result(
            DataToolOption=DataToolOption,
            ToolOptionFrame=ToolOptionFrame,
            GetToolOption=GetToolOption
        )"""

    def safe_item(df, column):
        if column not in df.columns:
            return NO_DATA

        value = df.item(0, column)

        if value is None:
            return NO_DATA

        return value

    ModelName = safe_item(ToolOptionFrame, "modelName")
    ModelSuffix = safe_item(ToolOptionFrame, "modelSuffix")
    appliedYear = safe_item(ToolOptionFrame, "appliedYear")
    tool1 = safe_item(ToolOptionFrame, "toolOption1")
    tool2 = safe_item(matchesSuffix, "toolOption2")
    tool3 = safe_item(ToolOptionFrame, "toolOption3")
    tool4 = safe_item(ToolOptionFrame, "toolOption4")
    tool5 = safe_item(ToolOptionFrame, "toolOption5")
    tool6 = safe_item(ToolOptionFrame, "toolOption6")
    tool7 = safe_item(ToolOptionFrame, "toolOption7")

    commercialTool = safe_item(ToolOptionFrame, "commercialToolOption")
    crcSum = safe_item(ToolOptionFrame, "crcChecksum")
    country = safe_item(ToolOptionFrame, "country")
    countryGroup = safe_item(ToolOptionFrame, "countryGroup")
    areaCode = safe_item(ToolOptionFrame, "areaCodeOrOption")
    boardOption = safe_item(ToolOptionFrame, "boardOption")

    return master_tool_option({
        "DataToolOption": DataToolOption,
        "ToolOptionFrame": ToolOptionFrame,
        "ModelName": ModelName,
        "ModelSuffix": ModelSuffix,
        "AppliedYear": appliedYear,
        "tool1": tool1,
        "tool2": tool2,
        "tool3": tool3,
        "tool4": tool4,
        "tool5": tool5,
        "tool6": tool6,
        "tool7": tool7,
        "commercialTool": commercialTool,
        "crcSum": crcSum,
        "country": country,
        "countryGroup": countryGroup,
        "areaCode": areaCode,
        "boardOption": boardOption
    })

    #return matchesSuffix, GetSuffixName


def getToolinfo(tool1, tool2, tool3, tool4, tool5, tool6, boardopt: str):
    target_boardoption = boardopt
    tooloption1 = tool1
    tooloption2 = tool2
    tooloption3 = tool3
    tooloption4 = tool4
    tooloption5 = tool5
    tooloption6 = tool6
    boardoption = int(target_boardoption,16)
  
    #GetBoardcheck = calculate_boardoption()
    GetTool1check = calculate_tooloption1()
    GetTool2check = calculate_tooloption2()
    GetTool3check = calculate_tooloption3()
    GetTool4check = calculate_tooloption4()
    GetTool5check = calculate_tooloption5()
    GetTool6check = calculate_tooloption6()
    GetBoardcheck = calculate_boardoption()

    GetTool1 = {
        "inch_code": inch_code,
        "tool_code": tool_code,
        "panel_code": panel_code,
        "pmic_mode": pmic_mode,
        "backlight_type": backlight_type,
        "localdin_block": localdim_block,
        "adjust_vcom": adjust_vcom
    }

    key_value = [
                {
                    "Inch Code": inch_code,
                    "Tool": tool_code,
                    "Panel Code": panel_code,
                    "Pmic set Mode": pmic_mode,
                    "Backlight Type": backlight_type,
                    "Local Dim Block": localdim_block,
                    "Adjust LCD VCOM": adjust_vcom,
                    "Audio Power Volt": audiopw_volt,
                    "PWM2 Duty": pwm2_duty,
                    "Max PWM": max_pwm,
                    "Scanning B/L": scanning_bl,
                    "BPL Current Type": bpl_type,
                    "Power Mode Select": powermode_select,
                    "Eye Sensor LED Gain Control": eyesensor_gain,
                    "Standard Backlight": standard_backlight,
                    "Vivid Sports Backlight": vividsports_bl,
                    "Panel Gamut Type": panelgamut_type,
                    "Motion Pro Mode": motionpro_mode,
                    "Support Dolby Vision HDR": support_dolbyvisionhdr,
                    "Audio Amp": ampchip_type,
                    "Speaker Watt": audiospk_watt,
                    "Luminance Up": support_luminanceup,
                    "Support VRR": support_vrr,
                    "Support Freesync": support_freesync,
                    "AI Sound": aisound_type,
                    "Dolby Vison Precision Detail": dolbyvison_precisiondetail,
                    "PWM Freq. double": pwmfreq_double,
                    "Motion Booster": motion_booster,
                    "OLED DPC Type": oleddpc_type,
                    "Support Freesync Pro": support_freesyncpro,
                    "Logo Light": logo_light,
                    "Fan Control": support_fancontrol,
                    "Camera Ready": camera_type,
                    "WOL": support_wol,
                    "HDMI Switch IC": hdmiswitch_type,
                    "Moving Speaker": moving_speaker,
                    "Support IR Blaster": support_irblaster,
                    "Mirror Mode": mirror_mode,
                    "Screen Roller Motor": screenroller_motor,
                    "Motion Detect Sensor": motiondetect_sensor,
                    "LCD Screen Cooling": lcdscreen_cooling,
                    "LAN Support": support_lan,
                    "Special Edition": special_edition,
                    "Support Optical Jack": optical_jack,
                    "Support Far Field Voice": farfield_voice,
                    "Type-C controller IC": type_cctric,
                    "AI Art Extension": aiart_extension,
                    "Board Design": board_design,
                    "EMMC Density": emmc_density,
                    "Video ADC Index": videoadc_index,
                    "Audio ADC Index": audioadc_index,
                    "USB Count": usb_count,
                    "USB 1": usb_1,
                    "USB 2": usb_2,
                    "USB 3": usb_3,
                    "Headphone": support_headphone,
                    "CI_AP Key": ciap_key,
                    "DTCP_AP Key": dtcpap_key,
                    "MAC_AP Key": macap_key,
                    "ESN_AP Key": esnap_key,
                    "Widevine_AP Key": widevineap_key,
                    "CI_ECP_AP Key": ciecpap_key,
                    "ATSC 3.0 (KR) Key": atsc30_key,
                    "New ErP": new_erp,
                    "OLED CPC": oled_cpc,
                    "PD Charge": pd_charge,
                    "Digital Eye": digital_eye,
                    "Energy Star": energy_star,
                    "AdaptiveDim Default": adaptivedim_default,
                    "Adaptive Dimming": adaptive_dimming,
                    "Default Backlight": default_backlight,
                    "Sound Mode": sound_mode,
                    "ISF": isf_mode,
                    "Set ID": set_id,
                    "Support Calman SW": support_calmansw,
                    "DVR Ready": dvr_ready,
                    "Audio Line Out": audio_lineout,
                    "WIFI/BT": wifi_bt,
                    "WIFI assy type": wifi_assy_type,
                    "Motion Remocon": motion_remocon,
                    "AudioEQ Eye": audioeq_eye,
                    "EDID Type": edid_type,
                    "Panel Type": panel_type,
                    "Panel Interface": panel_interface,
                    "Panel Resolution": panel_resolution,
                    "Panel Max Frame Rate": panel_maxframerate,
                    "Panel Max Bit": panel_maxbit,
                    "Panel Division Type": panel_divisiontype,
                    "External Frc Type": external_frc,
                    "DDR Size": ddr_size,
                    "Country Type": country_type,
                    "Tuner Type": tuner_type,
                    "Micom LocalKey Type": micom_localkey,
                    "AR RL on Control": arrlon_control,
                    "IR LED Control": irled_control,
                    "Model Type": model_type,
                    "Wee Turnkey Version": weeturnkey_version
                }
            ]

    df = pl.DataFrame(
        list(key_value[0].items()),
        schema=["Properties", "Target Model"]
    )

    return df



