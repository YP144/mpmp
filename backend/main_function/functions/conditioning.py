# PROPRIETARY AND CONFIDENTIAL
#
# This file and its contents are the proprietary property of LGERC HW ATSC 60Hz.
# Unauthorized copying, redistribution, modification, disclosure, or claim of
# ownership is strictly prohibited.
#
# Use of LLMs, AI agents, AI chatbots, or automated code tools to alter,
# redefine, disguise, or claim this source code as another party’s own work is
# prohibited and may result in legal action.
#
# Removal or deletion of this notice, whether intentional or unintentional,
# does not waive any rights and shall be considered a breach of the applicable
# repository license notice.
#
# ALL RIGHTS RESERVED.



#from streamlit_navigation_bar import st_navbar
import pandas as pd
import time
import requests
import json
import io
import numpy as np
from openpyxl import load_workbook, workbook
from openpyxl import Workbook 
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import polars as pl
from pathlib import Path
import os
from dej_sdk import bucket
import tempfile
import re
import polars as pl
import hashlib
from pandas.io.formats.style import Styler
import connectorx
from copy import copy as copy_style
import zipfile
import re
import threading
from collections import defaultdict


def zerobom_conditioning(modelFrame):       #for conditionig Zero State BOM json
    
    #with open(modelFrame, "r", encoding="utf-8-sig") as f:
    #    data = json.load(f)

    """if isinstance(modelFrame, (str, Path)):
        with open(modelFrame, "r", encoding="utf-8-sig") as f:
            data = json.load(f)

    elif isinstance(modelFrame, dict):
        data = modelFrame

    else:
        raise TypeError(
            "zerobom_conditioning() expects a file path or JSON dict."
        )"""

    if not isinstance(modelFrame, dict):
        raise TypeError(
            f"Expected API response as dict, got {type(modelFrame).__name__}"
        )

    if "data" not in modelFrame:
        raise ValueError(
            "Invalid BOM API response: missing 'data' key. "
            f"Available keys: {list(modelFrame.keys())}. "
            f"Response preview: {str(modelFrame)[:500]}"
        )
    

    all_outdata = []
    getSupplyType = []
    BOMSupplyType = modelFrame["data"]
    ModelName = BOMSupplyType[0].get("childNameDisplay")
    
    ModelName = ModelName.split("@")
    OnlyModelName = str(ModelName[0])

    for item in BOMSupplyType:
        all_outdata.append(item["outData"])
        getSupplyType.append(item["supplyType"])

    #Dataframejson = data['data']['outData']
    types_found = defaultdict(set)

    for row in all_outdata:
        for k, v in row.items():
            types_found[k].add(type(v).__name__)

    for col, types in sorted(types_found.items()):
        if len(types) > 1:
            print(f"PROBLEM COLUMN: {col}")
            print(f"TYPES: {types}")


    BOMjson = pl.DataFrame(all_outdata, infer_schema_length=None)
    BOMlist = pl.DataFrame(BOMSupplyType, infer_schema_length=None)
    #BOMSupplyType = pl.DataFrame(getSupplyType)
    BOMjson = BOMjson.select("pp_names", "pp_descriptions", "mr_siteTechnicalSpec", "mr_userItemType", "pp_revision", "currentChangeName", "pp_obid", "pp_states")
    SupplyBOM = BOMlist.select("bomLevel" ,"supplyType", "bomQuantity", "jobExplanation", "dataType", "bomExpandPath", "uniqueStr", "uniqueStrParent", "referenceDesignator")

    CompleteBOM = pl.concat([BOMjson, SupplyBOM], how="horizontal")
    #CompleteBOM = CompleteBOM.unique(subset=["pp_names"], keep="first")
    CompleteBOM = CompleteBOM.with_columns(
        pl.col("pp_descriptions").str.to_lowercase()
    )
    CompleteBOM = CompleteBOM.with_columns(
        pl.col("mr_siteTechnicalSpec").str.to_lowercase()
    )
    CompleteBOM = CompleteBOM.select("pp_names", "pp_descriptions", "mr_siteTechnicalSpec", "mr_userItemType", "supplyType", "bomQuantity", "jobExplanation", "pp_obid", "bomLevel", "dataType", "currentChangeName", "pp_revision", "uniqueStr", "uniqueStrParent", "referenceDesignator", "pp_states").rename({
            "pp_descriptions" : "part_descriptions",
            "pp_names" : "part_number",
            "supplyType" : "part_supplyType",
            "mr_userItemType": "part_UIT",
            "bomQuantity": "part_quantity",
            "mr_siteTechnicalSpec": "part_TechnicalSpec",
            "jobExplanation": "JobExpl",
            "pp_revision": "revisions",
            "currentChangeName": "changeNo"
        }
        
    )

    return CompleteBOM, OnlyModelName

def get_parent_partlist(CompleteBOM):
    create_parent_partnumber = CompleteBOM.select([
        "part_number",
        "uniqueStr",
        "bomLevel",
        "pp_obid"
    ]).rename({
        "part_number": "parent_part_number",
        "uniqueStr": "uniqueStrParent",
        "bomLevel": "parent_bomLevel"
    })

    parent_part_joined = CompleteBOM.join(
        create_parent_partnumber,
        on=["uniqueStrParent"],
        how="left"
    )

    # Get EBR parts
    get_ebr_part = parent_part_joined.filter(
        pl.col("part_number").fill_null("").str.starts_with("EBR") &
        (pl.col("bomLevel").cast(pl.Int64) == 5)
    )

    # Get EAX parts
    get_eax_part = parent_part_joined.filter(
        pl.col("part_number").fill_null("").str.starts_with("EAX")
    )

    # --------------------------------------------------
    # If EAX part/list is empty, return message dataframe
    # --------------------------------------------------
    if get_eax_part.is_empty():
        eax_not_detected_df = pl.DataFrame({
            "status": ["EAX not detected"]
        })

        return eax_not_detected_df, None, None

    # --------------------------------------------------
    # If EBR part/list is empty, also handle safely
    # --------------------------------------------------
    if get_ebr_part.is_empty():
        ebr_not_detected_df = pl.DataFrame({
            "status": ["EBR not detected"]
        })

        get_eax_num = get_eax_part.item(0, "part_number")

        return ebr_not_detected_df, get_eax_num, None

    get_eax_num = get_eax_part.item(0, "part_number")
    get_eax_obid = get_eax_part.item(0, "pp_obid")
    underEBR = get_ebr_part.item(0, "part_number")

    get_ebr_underEAX = parent_part_joined.filter(
        (pl.col("parent_part_number") == underEBR) &
        (pl.col("parent_bomLevel").cast(pl.Int64) == 5)
    )

    # --------------------------------------------------
    # If parts under EBR are empty
    # --------------------------------------------------
    if get_ebr_underEAX.is_empty():
        empty_under_ebr_df = pl.DataFrame({
            "status": ["No parts detected under EBR"]
        })

        return empty_under_ebr_df, get_eax_num, underEBR

    result = (
        get_ebr_underEAX.with_columns(
            pl.col("referenceDesignator")
            .str.split(",")
        )
        .explode("referenceDesignator")
    )

    return result, get_eax_num, underEBR, get_eax_obid
