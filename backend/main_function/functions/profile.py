# PROPRIETARY AND CONFIDENTIAL
#
# This file and its contents are the proprietary property of LGERC HW ATSC 60Hz.
# Unauthorized copying, redistribution, modification, disclosure, or claim of
# ownership is strictly prohibited.
#
# Use of LLMs, AI agents, AI chatbots, or automated code tools to alter,
# redefine, disguise, or claim this source code as another party’s own work is
# prohibited and may result in legal action.
#
# Removal or deletion of this notice, whether intentional or unintentional,
# does not waive any rights and shall be considered a breach of the applicable
# repository license notice.  
# 
# ALL RIGHTS RESERVED.  


#from streamlit_navigation_bar import st_navbar
import pandas as pd
import time
import requests
import json
import io
import numpy as np
from openpyxl import load_workbook, workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import polars as pl
from pathlib import Path
import os
from dej_sdk import bucket
import tempfile
import re
import hashlib
from pandas.io.formats.style import Styler
import connectorx
import traceback
from difflib import SequenceMatcher

from main_function.services.check_sso import SSO_NPDM
from main_function.services.main_hwpart import targetpart_compare, developmentpart_compare, initialpart_compare
from main_function.services.main_tooloption import run_tooloption, getToolinfo
from main_function.services.subs_hwpart import run_subshwpart
from datetime import timedelta, datetime, date

from main_function.functions.get_latestver import read_latest_version, read_table
from main_function.functions.getdata import (
    HerosURL, userURL, CoworkURL, memberURL,
    OPurl, suffixURL, extraURL, AccMaster
    )
from main_function.functions.get_profile_support import get_suffix_data
from main_function.functions.get_init_info import predict_development
from main_function.functions.datasend import send_to_pymongo_generic

from main_function.functions.profiling.get_year import db_search_model_with_fallback_score, search_model_with_fallback_score


#
# == obtaining name base on the regex pattern
#

def normalize_model_text(value: str) -> str:
    if value is None:
        return ""

    value = str(value).upper().strip()
    value = value.replace("_", "")
    value = re.sub(r"[^A-Z0-9]", "", value)

    return value

def extract_model_core(input_model: str) -> str:
    text = normalize_model_text(input_model)
    text = re.sub(r"^\d{2,3}(?=[A-Z])", "", text)
    return text

def similarity_score(a: str, b: str) -> float:
    if not a or not b:
        return 0.0

    return SequenceMatcher(None, a, b).ratio() * 100

def longest_common_substring_len(a: str, b: str) -> int:
    if not a or not b:
        return 0

    previous = [0] * (len(b) + 1)
    longest = 0

    for i in range(1, len(a) + 1):
        current = [0] * (len(b) + 1)

        for j in range(1, len(b) + 1):
            if a[i - 1] == b[j - 1]:
                current[j] = previous[j - 1] + 1
                longest = max(longest, current[j])

        previous = current

    return longest

def score_model_match(input_model: str, pattern: str) -> dict:
    input_clean = normalize_model_text(input_model)
    input_core = extract_model_core(input_model)
    pattern_clean = normalize_model_text(pattern)

    if not pattern_clean:
        return {
            "input_clean": input_clean,
            "input_core": input_core,
            "pattern_clean": pattern_clean,
            "is_exact_contains": False,
            "contains_score": 0.0,
            "similarity_percent": 0.0,
            "coverage_percent": 0.0,
            "lcs_length": 0,
            "pattern_length": 0,
            "final_score": 0.0,
        }

    is_exact_contains = pattern_clean in input_clean or pattern_clean in input_core

    contains_score = 100.0 if is_exact_contains else 0.0

    similarity_percent = max(
        similarity_score(input_clean, pattern_clean),
        similarity_score(input_core, pattern_clean)
    )

    lcs_length = max(
        longest_common_substring_len(input_clean, pattern_clean),
        longest_common_substring_len(input_core, pattern_clean)
    )

    coverage_percent = lcs_length / len(pattern_clean) * 100

    pattern_length = len(pattern_clean)

    length_bonus = min(pattern_length * 2, 20)

    final_score = (
        contains_score * 0.45
        + similarity_percent * 0.30
        + coverage_percent * 0.20
        + length_bonus * 0.05
    )

    if is_exact_contains:
        final_score += 10

    final_score = min(final_score, 100.0)

    return {
        "input_clean": input_clean,
        "input_core": input_core,
        "pattern_clean": pattern_clean,
        "is_exact_contains": is_exact_contains,
        "contains_score": round(contains_score, 2),
        "similarity_percent": round(similarity_percent, 2),
        "coverage_percent": round(coverage_percent, 2),
        "lcs_length": lcs_length,
        "pattern_length": pattern_length,
        "final_score": round(final_score, 2),
    }

def find_nearest_model_name_match(
    df: pl.DataFrame,
    input_model: str,
    model_col: str = "Model Name",
    top_n: int = 3,
    min_score: float = 0.0,
):
    """
    Find nearest matching row from model_col while keeping all original df columns.

    Returns:
        all_match_result:
            Original dataframe columns + scoring columns.

        show_top_3_matches:
            Top N matched rows with original columns + scoring columns.

        nearest_match:
            Best matched row only with original columns + scoring columns.
    """

    # Add row_index to preserve original row identity
    df_with_index = df.with_row_index("row_index")

    scored = []

    for row in df_with_index.select(["row_index", model_col]).iter_rows(named=True):
        row_index = row["row_index"]
        model_pattern = row[model_col]

        score_data = score_model_match(input_model, model_pattern)

        scored.append({
            "row_index": row_index,
            **score_data,
        })

    score_df = pl.DataFrame(scored)

    result = (
        df_with_index
        .join(score_df, on="row_index", how="left")
        .filter(pl.col("final_score") >= min_score)
        .sort(
            by=[
                "final_score",
                "is_exact_contains",
                "pattern_length",
                "coverage_percent",
                "similarity_percent",
            ],
            descending=[True, True, True, True, True],
        )
    )

    show_top_3_matches = result.head(top_n)
    nearest_match = result.head(1)

    return result, show_top_3_matches, nearest_match


#
# == filtering operation plan function ==
#

def _clean_value(value):
    if value is None:
        return ""

    value = str(value)

    value = value.replace("\u00a0", " ")
    value = value.replace("\ufeff", "")
    value = value.strip()

    if value.lower() in ["none", "null", "nan", "nat"]:
        return ""

    return value

def _is_blank(value):
    return _clean_value(value) == ""

def _split_rules(rule_cell):
    if _is_blank(rule_cell):
        return []

    rule_cell = _clean_value(rule_cell)

    return [
        rule.strip()
        for rule in re.split(r"[\/|]", rule_cell)
        if rule.strip()
    ]

def _rule_to_regex(rule):
    rule = _clean_value(rule)

    regex = re.escape(rule)

    # * means any characters
    regex = regex.replace(r"\*", ".*")

    # _ means any single character
    regex = regex.replace("_", ".")

    # DB sometimes uses . as any single character
    regex = regex.replace(r"\.", ".")

    return f"^{regex}$"

def _match_any_rule(value, rule_cell):
    rules = _split_rules(rule_cell)

    if not rules:
        return False

    value = _clean_value(value)

    if value == "":
        return False

    for rule in rules:
        pattern = _rule_to_regex(rule)

        if re.match(pattern, value, flags=re.IGNORECASE):
            return True

    return False

def _base_match(value, rule_cell):
    """
    Base column logic.

    If base cell is blank:
        means ANY, so return True.

    If base cell has rule:
        input value must match that rule.
    """
    if _is_blank(rule_cell):
        return True

    return _match_any_rule(value, rule_cell)

def _exception_hit(value, exception_cell):
    """
    Exception column logic.

    If exception cell is blank:
        means no exception, so return False.

    If exception cell has rule and input matches:
        exception is hit, so return True.
    """
    if _is_blank(exception_cell):
        return False

    return _match_any_rule(value, exception_cell)


def _part_number_match(part_number_prefix, row_part_number):
    part_number_prefix = _clean_value(part_number_prefix)
    row_part_number = _clean_value(row_part_number)

    if part_number_prefix == "":
        return True

    if row_part_number == "":
        return False

    pattern = rf"^{re.escape(part_number_prefix)}"

    return re.search(
        pattern,
        row_part_number,
        flags=re.IGNORECASE
    ) is not None
    
def parse_part_number_polars(
    rule_df: pl.DataFrame,
    part_number,
    model_name,
    model_plant,
    model_region,
    model_suffix,
    model_series,
    model_inch,
    part_number_column="partNo",
    only_valid=False
):
    scored_rows = []

    for idx, row in enumerate(rule_df.iter_rows(named=True)):

        part_number_match = _part_number_match(
            part_number,
            row.get(part_number_column)
        )

        # Optional: keep only rows where partNo starts with input part_number
        if not part_number_match:
            continue

        suffix_match = _base_match(
            model_suffix,
            row.get("suffixHq")
        )
        suffix_exception = _exception_hit(
            model_suffix,
            row.get("suffixHqException")
        )

        plant_match = _base_match(
            model_plant,
            row.get("plant")
        )
        plant_exception = _exception_hit(
            model_plant,
            row.get("plantException")
        )

        region_match = _base_match(
            model_region,
            row.get("region")
        )
        region_exception = _exception_hit(
            model_region,
            row.get("regionException")
        )

        model_name_match = _base_match(
            model_name,
            row.get("model")
        )
        model_name_exception = _exception_hit(
            model_name,
            row.get("modelException")
        )

        series_match = _base_match(
            model_series, 
            row.get("operationSeries")
        )
        series_exception = _exception_hit(
            model_series,
            row.get("operationSeriesException")
        )

        inch_match = _base_match(
            model_inch,
            row.get("operationInch")
        )
        inch_exception = _exception_hit(
            model_inch,
            row.get("operationInchException")
        )

        is_valid = (
            part_number_match
            and suffix_match
            and not suffix_exception
            and plant_match
            and not plant_exception
            and region_match
            and not region_exception
            and model_name_match
            and not model_name_exception
            and series_match
            and not series_exception
            and inch_match
            and not inch_exception
        )

        scores = {
            "part_number_score": int(part_number_match),

            "suffix_score": int(suffix_match),
            "suffix_exception_score": int(not suffix_exception),

            "plant_score": int(plant_match),
            "plant_exception_score": int(not plant_exception),

            "region_score": int(region_match),
            "region_exception_score": int(not region_exception),

            "model_name_score": int(model_name_match),
            "model_name_exception_score": int(not model_name_exception),

            "series_score": int(series_match),
            "series_exception_score": int(not series_exception),

            "inch_score": int(inch_match),
            "inch_exception_score": int(not inch_exception),
        }

        match_score = sum(scores.values())
        max_score = len(scores)
        match_percent = round((match_score / max_score) * 100, 2)

        failed_reasons = []

        if not part_number_match:
            failed_reasons.append("part_number_not_match")

        if not suffix_match:
            failed_reasons.append("suffix_not_match")
        if suffix_exception:
            failed_reasons.append("suffix_exception_hit")

        if not plant_match:
            failed_reasons.append("plant_not_match")
        if plant_exception:
            failed_reasons.append("plant_exception_hit")

        if not region_match:
            failed_reasons.append("region_not_match")
        if region_exception:
            failed_reasons.append("region_exception_hit")

        if not model_name_match:
            failed_reasons.append("model_name_not_match")
        if model_name_exception:
            failed_reasons.append("model_name_exception_hit")

        if not series_match:
            failed_reasons.append("series_not_match")
        if series_exception:
            failed_reasons.append("series_exception_hit")

        if not inch_match:
            failed_reasons.append("inch_not_match")
        if inch_exception:
            failed_reasons.append("inch_exception_hit")

        raw_row = {
            f"raw_{k}": _clean_value(v)
            for k, v in row.items()
        }

        output_row = {
            "row_index": idx,

            "is_valid": bool(is_valid),
            "match_score": int(match_score),
            "max_score": int(max_score),
            "match_percent": float(match_percent),
            "failed_reasons": ", ".join(failed_reasons),

            "input_part_number": _clean_value(part_number),
            "row_part_number": _clean_value(row.get(part_number_column)),
            "part_number_match": bool(part_number_match),

            "input_model_name": _clean_value(model_name),
            "input_model_plant": _clean_value(model_plant),
            "input_model_region": _clean_value(model_region),
            "input_model_suffix": _clean_value(model_suffix),
            "input_model_series": _clean_value(model_series),
            "input_model_inch": _clean_value(model_inch),

            "suffix_match": bool(suffix_match),
            "suffix_exception": bool(suffix_exception),

            "plant_match": bool(plant_match),
            "plant_exception": bool(plant_exception),

            "region_match": bool(region_match),
            "region_exception": bool(region_exception),

            "model_name_match": bool(model_name_match),
            "model_name_exception": bool(model_name_exception),

            "series_match": bool(series_match),
            "series_exception": bool(series_exception),

            "inch_match": bool(inch_match),
            "inch_exception": bool(inch_exception),

            **scores,
            **raw_row,
        }

        if only_valid:
            if is_valid:
                scored_rows.append(output_row)
        else:
            scored_rows.append(output_row)

    if not scored_rows:
        empty_df = pl.DataFrame()
        return empty_df, empty_df

    all_result = pl.DataFrame(scored_rows, infer_schema_length=None)

    matched_100_result = all_result.filter(
        pl.col("match_percent") == 100
    )

    return all_result, matched_100_result

def get_model_year(self, model_name):

        get_result_prefix = ""

        get_latest_26Y = read_latest_version(
            db_uri=HerosURL,
            prefix="HEROmodel2026Y_v3.3_v"
        )

        get_latest_27Y = read_latest_version(
            db_uri=HerosURL,
            prefix="HEROmodel2027Y_v1.0_v"
        )

        result_year = search_model_with_fallback_score(
            model_input=model_name,
            df_26y=get_latest_26Y,
            df_27y=get_latest_27Y,            
        )

        result_year_location = result_year["matched_location"]

        if result_year_location == "26Y":
            get_result_prefix = "HEROmodel2026Y_v3.3_v"
        elif result_year_location == "27Y":
            get_result_prefix = "HEROmodel2027Y_v1.0_v"

        return get_result_prefix

class ModelProfile():
    def __init__(self, targetModel: str, ssolgenet_cookie: str = "", smsession_cookie: str = ""):
        self.targetModel = targetModel
        self.ssolgenet_cookie = ssolgenet_cookie
        self.smsession_cookie = smsession_cookie
        
        self.getProfile()

    def get_model_year(self, model_name):

        get_result_prefix = ""

        get_latest_26Y = read_latest_version(
            db_uri=HerosURL,
            prefix="HEROmodel2026Y_v3.3_v"
        )

        get_latest_27Y = read_latest_version(
            db_uri=HerosURL,
            prefix="HEROmodel2027Y_v1.0_v"
        )

        result_year = search_model_with_fallback_score(
            model_input=model_name,
            df_26y=get_latest_26Y,
            df_27y=get_latest_27Y,            
        )

        result_year_location = result_year["matched_location"]

        if result_year_location == "26Y":
            get_result_prefix = "HEROmodel2026Y_v3.3_v"
        elif result_year_location == "27Y":
            get_result_prefix = "HEROmodel2027Y_v1.0_v"

        return get_result_prefix


    def getProfile(self):
        # Profile Matching
        sso = SSO_NPDM()
        # Use instance attributes or fallback gracefully
        is_logged_in = sso.auto_login(
            ssolgenet_cookie=self.ssolgenet_cookie, 
            smsession_cookie=self.smsession_cookie
        )

        name_user = sso.userIc
        name_ic = sso.userId

        user_row = None
        get_year_prefix = self.get_model_year(self.targetModel)
        #Check engineer profile
        user_row = read_table(
            memberURL,
            "user_list"
        )

        get_name_row = user_row.filter(
            pl.col("ID").str.contains(name_user)
        )

        if get_name_row.height > 0:
            self.user_team = get_name_row.item(0, "TEAM")
            self.user_part = get_name_row.item(0, "Part")
            self.user_level = get_name_row.item(0, "Level")
        else:
            self.user_team = "No TEAM"
            self.user_part = "NO PART"
            self.user_level = "NO LEVEL"

        #get_dev_model = predict_development(self.targetModel)
    
        GetModelName = self.targetModel.split(".")
        self.firstModelName = GetModelName[0]

        query = "SELECT * FROM public.all_engineer"
        getUser = pl.read_database_uri(
            uri=userURL,
            query=query
            #engine="connectorx"
        )

        getUser = getUser.select("getName") #GetData list with "getName" as headers

        gettoday = date.today() #Get Todays Date

        pastday = gettoday - timedelta(days=365) 

        cowork_params = {
        "cParamSearchVO":{
            "releaseCreationDate":{
                "comparison":"BT",
                "fromDate":"",
                "toDate":""
            },
            "gerpCreationDate":{
                "fromDate":f"{pastday}",
                "toDate":f"{gettoday}",
                "comparison":"BT"
            },
            "requestBOMOpen":{
                "comparison":"BT",
                "fromDate":"",
                "toDate":""
            },
            "divisions":"GLZ",
            "creationPlantUnit":"GLZ.EKHQ",
            "names":f"{self.targetModel}"
            }
        }

        cowork_headers = {
            "Accept": "application/json, text/plain, */*",
            "Authorization": sso.auth_bearer,
            "BusinessUnit": "HE",
            "Content-Type": "application/json",
            "Host": "npdm.lge.com",
            "Referer": "https://npdm.lge.com/"
        }

        cowork_response = requests.put(CoworkURL, headers=cowork_headers, json=cowork_params)
        getjson_cowork = cowork_response.json()

        try:
            if (
                getjson_cowork is not None
                and "data" in getjson_cowork
                and getjson_cowork["data"] is not None
                and len(getjson_cowork["data"]) > 0
            ):
                all_outdata = []
                coworkdata = getjson_cowork["data"]

                for item in coworkdata:
                    if item.get("outData") is not None:
                        all_outdata.append(item["outData"])

                required_outdata_columns = [
                    "cha_coWorker1", "cha_coWorker2", "cha_coWorker3",
                    "cha_coWorker4", "cha_coWorker5", "cha_coWorker6",
                    "cha_coWorker7", "cha_coWorker8", "cha_coWorker9",
                    "cha_names", "cha_states", "mr_chassis"
                ]

                required_cowork_columns = [
                    "names", "titles", "technicalSpecification", "states"
                ]

                if len(all_outdata) == 0:
                    self.getCircuit = "Not Found"

                else:
                    getcoworkName = pl.DataFrame(all_outdata)

                    for col in required_outdata_columns:
                        if col not in getcoworkName.columns:
                            getcoworkName = getcoworkName.with_columns(
                                pl.lit("").alias(col)
                            )

                    getcoworkName = getcoworkName.select(required_outdata_columns)

                    coworkdata = pl.DataFrame(coworkdata, infer_schema_length=1000)

                    for col in required_cowork_columns:
                        if col not in coworkdata.columns:
                            coworkdata = coworkdata.with_columns(
                                pl.lit("").alias(col)
                            )

                    coworkdata = coworkdata.select(required_cowork_columns)

                    Result_cowork = pl.concat([getcoworkName, coworkdata], how="horizontal")

                    atsc_members = (
                        getUser
                        .select("getName")
                        .to_series()
                        .drop_nulls()
                        .cast(pl.Utf8)
                        .str.strip_chars()
                        .to_list()
                    )

                    search_columns = [
                        "cha_coWorker1", "cha_coWorker2", "cha_coWorker3",
                        "cha_coWorker4", "cha_coWorker5", "cha_coWorker6",
                        "cha_coWorker7", "cha_coWorker8", "cha_coWorker9",
                        "cha_names",
                        "names"
                    ]

                    def find_atsc_members(text):
                        if text is None:
                            return ""

                        found = []

                        for member in atsc_members:
                            if member in text:
                                found.append(member)

                        return ", ".join(found)

                    Result_atsc = Result_cowork.with_columns(
                        pl.concat_str(
                            [
                                pl.col(col).cast(pl.Utf8).fill_null("")
                                for col in search_columns
                            ],
                            separator=" "
                        ).alias("all_names")
                    ).with_columns(
                        pl.col("all_names")
                        .map_elements(find_atsc_members, return_dtype=pl.Utf8)
                        .alias("found_atsc_members")
                    ).filter(
                        pl.col("found_atsc_members") != ""
                    )

                    if Result_atsc.height > 0:
                        self.getCircuit = Result_atsc.item(0, "found_atsc_members")
                    else:
                        self.getCircuit = "Not Found"

            else:
                self.getCircuit = "Not Found"

        except Exception:
            self.getCircuit = "Not Found"
            
        #Check Year - Determine the Model Year
        model_year_crosschecked = db_search_model_with_fallback_score(self.firstModelName)

        #Get Model Year Identification
        self.model_year = model_year_crosschecked["matched_location"]
        
        #Get HEROs Row
        self.HEROsFrame = read_latest_version(
            db_uri=HerosURL,
            prefix=get_year_prefix
        )

        #Get Acc HEROs row
        self.getAccpart = read_latest_version(
            db_uri=AccMaster,
            prefix="HEROacc2026Y_v3.1_v"
        )


        #Identify by Suffix Name
        RegionFrame = read_table(
            db_uri=suffixURL,
            table_name="regionSuffix"
        )

        PlantFrame = read_table(
            db_uri=suffixURL,
            table_name="plantSuffix"
        )

        BoardFrame = read_table(
            db_uri=suffixURL,
            table_name="boardSuffix"
        )

        PlantFrame = PlantFrame.with_columns(
            pl.col("compareString").str.replace_all(r"_",".")
        )

        if len(GetModelName) != 2:
            raise ValueError(f"Invalid model format: {self.targetModel}")

        GetName, self.GetSuffixName = GetModelName
        self.GetModelName = GetName
        self.GetSuffixName = self.GetSuffixName.lower()

        self.GetRegionSuffix = self.GetSuffixName[1:3]
        self.GetBoardSuffix = self.GetSuffixName[4]
        BoardFrame = BoardFrame.select(["compareString", "Brand"]).with_columns(
            pl.col("compareString").str.to_lowercase().str.replace_all(r"_",".")
        )
        BoardFrame = BoardFrame.filter(
            pl.col("compareString").str.contains(self.GetBoardSuffix)
        )
        self.GetBoardInfo = BoardFrame.item(0, "Brand")

        GetPanel = self.GetSuffixName[3]
        SearchRegion = RegionFrame.filter(
            pl.col("compareString")
            .fill_null("")
            .str.to_lowercase()
            .str.contains(self.GetRegionSuffix.lower(), literal=True)
        )

        if SearchRegion.height == 0:
            raise ValueError(f"No region found for {self.GetRegionSuffix}")

        self.ParseRegion = SearchRegion[0, "Region"].upper()
        self.ParseMarket = SearchRegion[0, "Market"]

        #Get Plant from Database
        GetPlant = PlantFrame.filter(
            pl.lit(self.GetSuffixName.upper())
            .str.contains("^" + pl.col("compareString") + "$")
        )

        self.GetSubs = GetPlant[0, "Plant"]
        
        #Get HEROs data
        input_key = GetName.lstrip("0123456789").strip("_")
        
        # 1. Filter by Region and Panel (Suffix) to narrow down the search
        filtered_heros = self.HEROsFrame.with_columns(
            pl.col("Region").fill_null("").str.to_uppercase().str.strip_chars()
        ).filter(
            (pl.col("Region") == self.ParseRegion) &
            (pl.col("Suffix").fill_null("").str.to_uppercase().str.contains(GetPanel.upper(), literal=True))
        )

        # 2. Use your existing fuzzy match function to find the best model name
        if filtered_heros.height > 0:
            _, _, self.GetProfileRow = find_nearest_model_name_match(
                df=filtered_heros,
                input_model=GetName,
                model_col="Model Name",
                top_n=1,
                min_score=40
            )
        else:
            # Fallback: If Region/Panel filter is too strict, search the whole HEROs database
            _, _, self.GetProfileRow = find_nearest_model_name_match(
                df=self.HEROsFrame,
                input_model=GetName,
                model_col="Model Name",
                top_n=1,
                min_score=40
            )

        #Get Remote Wifi, Wireless, and MOdelName, SoC, Market
        if self.GetProfileRow.height > 0:
            acc_region = self.GetProfileRow[0, "Region"]
        else:
            acc_region = self.ParseRegion

        self.acc_model = (
            self.getAccpart
            .with_columns([
                pl.col("Model Name")
                .str.replace_all(r"^_+|_+$", "")
                .str.replace(r"^\d+", "")
                .alias("model_key")
            ])
            .filter(
                (pl.col("Region") == acc_region) &
                (pl.lit(input_key).str.starts_with(pl.col("model_key")))
            )
        )

        # Conditioning ACC -> Master List (MR26GA -> MR26GC)

        def normalize_remote_magic(remote_magic, suffix_name):
            """
            Conditioning ACC -> Master List

            Rule:
                MR26GA -> MR26GC

            Exception:
                If suffix_name[1:3] == "ff", do not rewrite.

            Examples:
                MR26GA            -> MR26GC
                MR26GA Metallic  -> MR26GC Metallic
                MR26GC            -> MR26GC
                MR26GC Metallic  -> MR26GC Metallic
            """

            if remote_magic is None:
                return remote_magic

            remote_magic = str(remote_magic).strip()

            # Do not rewrite if self.GetSuffixName[1:3] is "ff"
            if str(suffix_name)[1:3].lower() == "ff":
                return remote_magic

            # Rewrite only exact MR26GA word/code
            remote_magic = re.sub(
                r"\bMR26GA\b",
                "MR26GC",
                remote_magic,
                flags=re.IGNORECASE
            )

            return remote_magic

        self.SuffixCountry = self.GetSuffixName[1:3]

        # Conditioning ACC -> Master List (MR26GA -> MR26GC)
        RemoteMagic = self.acc_model[0, "Magic Remote"]

        RemoteMagic = normalize_remote_magic(
            remote_magic=RemoteMagic,
            suffix_name=self.SuffixCountry
        )

        self.ModelProfile = self.GetProfileRow
        self.HEROsFrame = self.HEROsFrame
        self.GetInch = self.GetProfileRow[0, "Inch"]
        self.GetSeries = self.GetProfileRow[0, "Series"]
        self.GetCPunit = self.GetProfileRow[0, "CP Unit"]
        self.GetGrade = self.GetProfileRow[0, "Grade"]
        self.GetSoC = self.GetProfileRow[0, "SoC"]
        self.GetChassis = self.GetProfileRow[0, "Chassis Name"]
        #self.GetPL = GetProfileRow[0, "PIC"]
        self.GetPanelName = self.GetProfileRow[0, "Panel Maker"]
        self.GetRemoteMR = RemoteMagic
        self.GetRemoteSR = self.acc_model[0, "Wifi Name"]
        #self.GetBoardType = self.GetProfileRow[0, "CM00886"]
        self.GetHerosRegion = self.GetProfileRow[0, "Region"]

        #send_to_pymongo_generic(model_year_crosschecked, "modelYearCheck")

        self.jsonProfile = [{
                "modelYear" : self.model_year,
                "ModelName" : self.targetModel,
                "Inch" : self.GetInch,
                "Series" : self.GetSeries,
                "Region": self.GetHerosRegion,
                "CPunit": self.GetCPunit,
                "SoC": self.GetSoC,
                "Board": self.GetBoardInfo,
                "Chassis": self.GetChassis,
                "Panel": self.GetPanelName,
                "RemoteMR": self.GetRemoteMR,
                "RemoteSR": self.GetRemoteSR,
                "Plant": self.GetSubs,
                "CircuitPIC": self.getCircuit,
                "ParseRegion": self.ParseRegion,
                "Market": self.ParseMarket,
                "ModelSuffix": self.SuffixCountry,
                "Model_Name": GetName,
                "user-info": {
                    "retest-userIc": name_ic,
                    "retest-userId": name_user,
                    "user-team": self.user_team,
                    "user-part": self.user_part,
                    "user-level": self.user_level
                }
            }
        ]




class GetOperationPart:
    def __init__(self, targetModel: str):
        self.targetModel = targetModel
        self.mongo_notification = {}
        self.unmatched_input_parts = []
        self.multiple_matches_report = []
        self.duplicated_parts_report = []
        self.opertionPlanMacthReport = {}
        self.getOPpart()

    def get_model_year(self, model_name):
        get_latest_26Y = read_latest_version(db_uri=HerosURL, prefix="HEROmodel2026Y_v3.3_v")
        get_latest_27Y = read_latest_version(db_uri=HerosURL, prefix="HEROmodel2027Y_v1.0_v")

        result_year = search_model_with_fallback_score(
            model_input=model_name,
            df_26y=get_latest_26Y,
            df_27y=get_latest_27Y,            
        )

        result_year_location = result_year["matched_location"]

        if result_year_location == "26Y":
            return "2026YHWplan_v", "HEROacc2026Y_v4.3_v"
        elif result_year_location == "27Y":
            return "2027YHWplan_v", "HEROSacc2027_v3.1_v"

        return "2026YHWplan_v", "HEROacc2026Y_v4.3_v"

    def getOPpart(self):
        profile = ModelProfile(self.targetModel)
        get_year_prefix, get_acc_prefix = self.get_model_year(self.targetModel)

        model_plant = profile.GetSubs
        model_region = profile.GetRegionSuffix
        model_suffix = profile.GetSuffixName
        model_series = profile.GetSeries
        model_inch = profile.GetInch

        def get_operation_plan_list(self_obj):
            get_latest_op_table = read_latest_version(db_uri=OPurl, prefix=get_year_prefix)
            get_part_number = get_latest_op_table.select(["partNo"]).with_columns(
                pl.col("partNo").str.slice(0, 3).alias("partNo_prefix")
            )
            unique_prefixes_list = get_part_number.select(["partNo_prefix"]).unique()["partNo_prefix"].to_list()
            return unique_prefixes_list, get_latest_op_table

        get_unique_op_list, GetOPlist = get_operation_plan_list(self)

        GetOPlist = GetOPlist.with_columns(pl.col("suffixHq").str.to_lowercase())
        self.region_operationPlan = profile.GetRegionSuffix

        GetAccMaster = read_latest_version(db_uri=AccMaster, prefix=get_acc_prefix)
        filtered_AccMaster = GetAccMaster.filter(pl.col("Region") == profile.GetHerosRegion)
        
        if filtered_AccMaster.height > 0:
            GetAccMaster = filtered_AccMaster

        if "row_index" in GetAccMaster.columns:
            GetAccMaster = GetAccMaster.drop("row_index")

        firstModelName = profile.firstModelName

        all_match_result, show_top_3_matches, nearest_match = find_nearest_model_name_match(
            df=GetAccMaster,
            input_model=firstModelName,
            model_col="Model Name",
            top_n=3,
            min_score=40
        )

        model_remote = nearest_match.item(0, "Magic Remote") if isinstance(nearest_match, pl.DataFrame) and nearest_match.height > 0 else ""

        # Normalize model_remote check for Ready / Blank / None state
        normalized_remote_val = str(model_remote).strip().lower() if model_remote else ""
        should_skip_remote_and_battery = normalized_remote_val in ["ready", "blank", "none", ""]

        all_op_dataframes = []
        all_unmatched_dataframes = [] 
        dynamic_match_reports = {}

        category_map = {
            "AKB": "Remote", "EAB": "Speaker", "EAD": "PowerCord",
            "ECC": "Edid", "EBS": "HeatPlate", "EAT": "WiFi", "EAC": "Battery", "691": "Battery"
        }

        for prefix in get_unique_op_list:
            search_prefix = prefix
            category_name = category_map.get(prefix, f"Category_{prefix}")

            # Check if this prefix belongs to Remote or Battery
            if should_skip_remote_and_battery and category_name in ["Remote", "Battery"]:
                # Record in report that this component is skipped due to model_remote state
                self.unmatched_input_parts.append({
                    "input_part_number": search_prefix,
                    "prefix": prefix,
                    "category": category_name,
                    "status": f"Skipped (model_remote is '{model_remote}')"
                })
                continue

            judge_result, op_df = parse_part_number_polars(
                rule_df=GetOPlist,
                part_number=search_prefix,
                model_name=profile.targetModel,
                model_plant=model_plant,
                model_region=model_region.upper(),
                model_suffix=model_suffix.upper(),
                model_series=model_series,
                model_inch=model_inch,
                part_number_column="partNo",
                only_valid=False
            )

            common_cols = list(set(judge_result.columns).intersection(set(op_df.columns)))
            unmatched_df = judge_result.join(op_df, on=common_cols, how="anti")

            base_criteria_str = (
                f"Part: {search_prefix} | Model: {profile.targetModel} | "
                f"Plant: {model_plant} | Region: {model_region.upper()} | "
                f"Suffix: {model_suffix.upper()} | Series: {model_series} | Inch: {model_inch}"
            )

            op_df = op_df.with_columns(
                pl.lit(f"Successfully matched base criteria for {search_prefix}").alias("reason"),
                pl.lit(base_criteria_str).alias("matched_criteria")
            )
            unmatched_df = unmatched_df.with_columns(
                pl.lit(f"Failed model/region/suffix/plant criteria for {search_prefix}").alias("reason"),
                pl.lit("Failed to match base model specifications").alias("matched_criteria")
            )

            if prefix == "AKB":
                remote_code = model_remote.replace("Metallic", "").strip() if model_remote else ""
                is_metallic = "Metallic" in model_remote if model_remote else False

                if remote_code:
                    base_filter = (
                        pl.col("raw_masterRemark")
                        .fill_null("")
                        .cast(pl.Utf8)
                        .str.strip_chars()
                        .str.contains(rf"(?i)\b{re.escape(remote_code)}\b")
                    )
                    akb_condition = base_filter & pl.col("raw_masterRemark").fill_null("").cast(pl.Utf8).str.contains(r"(?i)\bMetallic\b") if is_metallic else base_filter & ~pl.col("raw_masterRemark").fill_null("").cast(pl.Utf8).str.contains(r"(?i)\bMetallic\b")
                else:
                    akb_condition = pl.lit(True)

                akb_criteria_str = base_criteria_str + f" | Remote Code: {remote_code} | Metallic: {is_metallic}"

                akb_matched = op_df.filter(akb_condition).with_columns(
                    pl.lit(f"Successfully matched criteria and Remote/Metallic check for {search_prefix}").alias("reason"),
                    pl.lit(akb_criteria_str).alias("matched_criteria")
                )
                akb_unmatched = op_df.filter(~akb_condition).with_columns(
                    pl.lit(f"Failed Remote/Metallic check (Expected: {remote_code}, Metallic: {is_metallic})").alias("reason"),
                    pl.lit(base_criteria_str + " (Failed Remote/Metallic check)").alias("matched_criteria")
                )

                op_df = akb_matched
                unmatched_df = pl.concat([unmatched_df, akb_unmatched], how="diagonal")

            op_df = op_df.with_columns(pl.lit(category_name).alias("Category"))
            unmatched_df = unmatched_df.with_columns(pl.lit(category_name).alias("Category"))

            matched_row_count = op_df.height

            if matched_row_count == 0:
                self.unmatched_input_parts.append({
                    "input_part_number": search_prefix,
                    "prefix": prefix,
                    "category": category_name,
                    "status": "No Match Found"
                })

            if matched_row_count > 1:
                self.multiple_matches_report.append({
                    "input_part_number": search_prefix,
                    "prefix": prefix,
                    "category": category_name,
                    "matched_rows_count": matched_row_count,
                    "detected_rows": op_df.to_dicts()
                })

            if "partNo" in op_df.columns:
                part_no_counts = op_df.group_by("partNo").len().filter(pl.col("len") > 1)
                if part_no_counts.height > 0:
                    dup_part_numbers = part_no_counts["partNo"].to_list()
                    duplicated_rows = op_df.filter(pl.col("partNo").is_in(dup_part_numbers))
                    self.duplicated_parts_report.append({
                        "prefix": prefix,
                        "category": category_name,
                        "duplicated_part_numbers": dup_part_numbers,
                        "duplicated_count": duplicated_rows.height,
                        "duplicated_records": duplicated_rows.to_dicts()
                    })

            all_op_dataframes.append(op_df)
            all_unmatched_dataframes.append(unmatched_df)

            dynamic_match_reports[category_name] = {
                "partNo": search_prefix,
                "matchedCount": matched_row_count,
                "AllmatchReport": judge_result.to_dicts(),
                "matchReport": op_df.to_dicts(),
                "unmatchReport": unmatched_df.to_dicts()
            }

        self.OperationGet = pl.concat(all_op_dataframes, how="diagonal").with_row_index(name="row_number") if all_op_dataframes else pl.DataFrame()
        self.OperationUnmatched = pl.concat(all_unmatched_dataframes, how="diagonal").with_row_index(name="row_number") if all_unmatched_dataframes else pl.DataFrame()

        # Construction of initial base report struct
        self.opertionPlanMacthReport = {
            "timestamp": datetime.now(),
            "miseEnPlace": {
                "modelName": profile.targetModel,
                "model_name": firstModelName,
                "suffixName": model_suffix,
                "regionName": model_region,
                "herosRegion": profile.GetHerosRegion,
                "model_remote": model_remote,
                "profile": profile.jsonProfile,
                "accChosen": (
                    nearest_match.to_dicts()
                    if isinstance(nearest_match, pl.DataFrame)
                    else [{"nearest_match": str(nearest_match)}]
                ),
                "missingRequiredPartsSummary": "Pending Evaluation...",
                "top5UnmatchedCandidates": []
            },
            "unmatchedInputParts": self.unmatched_input_parts,
            "multipleMatchesReport": self.multiple_matches_report,
            "duplicatedPartsReport": self.duplicated_parts_report
        }

        self.opertionPlanMacthReport.update(dynamic_match_reports)

    def finalize_and_push_report(self, missing_summary: str, missing_details: list):
        """
        Updates missing required parts summary and top 5 match candidates directly into miseEnPlace,
        and pushes the final document to MongoDB.
        """
        top_candidates = []

        if self.OperationUnmatched.height > 0 and len(missing_details) > 0:
            for missing_item in missing_details:
                target_prefix = missing_item.get("prefix", "").split(" ")[0]
                
                # Filter unmatched dataframe for rows belonging to this missing prefix/category
                unmatched_sub = self.OperationUnmatched.filter(
                    pl.col("partNo").str.slice(0, 3) == target_prefix
                ) if "partNo" in self.OperationUnmatched.columns else pl.DataFrame()

                if unmatched_sub.height > 0:
                    score_col = "match_percent" if "match_percent" in unmatched_sub.columns else ("score" if "score" in unmatched_sub.columns else None)
                    
                    if score_col:
                        top_5 = unmatched_sub.sort(score_col, descending=True).head(5)
                    else:
                        top_5 = unmatched_sub.head(5)

                    top_candidates.append({
                        "prefix": target_prefix,
                        "category": missing_item.get("category"),
                        "top_5_matches": top_5.to_dicts()
                    })

        self.opertionPlanMacthReport["miseEnPlace"]["missingRequiredPartsSummary"] = missing_summary
        self.opertionPlanMacthReport["miseEnPlace"]["missingRequiredPartsDetails"] = missing_details
        self.opertionPlanMacthReport["miseEnPlace"]["top5UnmatchedCandidates"] = top_candidates

        # Send clean safe report to PyMongo
        safe_report = make_mongo_safe(self.opertionPlanMacthReport)
        self.mongo_notification = send_to_pymongo_generic(safe_report, "operationMatch")
        self.opertionPlanMacthReport["mongoNotification"] = self.mongo_notification