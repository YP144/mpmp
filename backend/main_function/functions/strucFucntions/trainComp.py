# PROPRIETARY AND CONFIDENTIAL
#
# This file and its contents are the proprietary property of LGERC HW ATSC 60Hz.
# Unauthorized copying, redistribution, modification, disclosure, or claim of
# ownership is strictly prohibited.
#
# Use of LLMs, AI agents, AI chatbots, or automated code tools to alter,
# redefine, disguise, or claim this source code as another party’s own work is
# prohibited and may result in legal action.
#
# Removal or deletion of this notice, whether intentional or unintentional,
# does not waive any rights and shall be considered a breach of the applicable
# repository license notice.
#
# ALL RIGHTS RESERVED.

from pathlib import Path
import sys

PROJECT_ROOT = Path(__file__).resolve().parents[2]

if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))



#from streamlit_navigation_bar import st_navbar
import pandas as pd
import time
import requests
import json
import io
import numpy as np
from openpyxl import load_workbook, workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import polars as pl
from pathlib import Path
import os
from dej_sdk import bucket
import tempfile
import re
import polars as pl
import hashlib
from pandas.io.formats.style import Styler
import connectorx
from datetime import datetime

from functions.conditioning import zerobom_conditioning
from functions.strucFucntions.quickCompare import (
#from .quickCompare import (
 
    GoGetTemplate,
    match_template_bom,
    build_match_expr,
    check_status,
    final_result_compare,
)

def get_structure_compare(target_model, bomFrame, ssolgenet_cookie=None, smsession_cookie=None):
    TemplateLinkUrl, GetTemplateFrame, GetTemplate = GoGetTemplate(target_model)
    joined_result_matched, run_info, output_report = match_template_bom(GetTemplateFrame, bomFrame, target_model)
    get_compare_result = check_status(joined_result_matched)
    joined_result_excel, json_data, compare_matching = final_result_compare(get_compare_result, GetTemplateFrame, run_info, output_report)

    result_check = (
        "NG"
        if joined_result_excel.select(pl.col("status").eq("NG").any()).item()
        else "OK"
    )

    return joined_result_excel, json_data, compare_matching, result_check

#st.write(PROJECT_ROOT)

#TargetBOM = os.path.abspath(f"{PROJECT_ROOT}/json/0_target_ebom.json")

#GetTargetBOM, TargetName = zerobom_conditioning(TargetBOM)

#joined_result_excel, json_data, compare_matching = get_structure_compare(TargetName, GetTargetBOM)

#st.dataframe(json_data)




