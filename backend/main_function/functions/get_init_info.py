# PROPRIETARY AND CONFIDENTIAL
#
# This file and its contents are the proprietary property of LGERC HW ATSC 60Hz.
# Unauthorized copying, redistribution, modification, disclosure, or claim of
# ownership is strictly prohibited.
#
# Use of LLMs, AI agents, AI chatbots, or automated code tools to alter,
# redefine, disguise, or claim this source code as another party’s own work is
# prohibited and may result in legal action.
#
# Removal or deletion of this notice, whether intentional or unintentional,
# does not waive any rights and shall be considered a breach of the applicable
# repository license notice.
#
# ALL RIGHTS RESERVED.

import pandas as pd
import time
import requests
import json
import io
import numpy as np
from openpyxl import load_workbook, workbook, Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import polars as pl
from pathlib import Path
import os
from dej_sdk import bucket
import tempfile
import re
import hashlib
from pandas.io.formats.style import Styler
import connectorx
from copy import copy as copy_style
import zipfile
import threading


def get_development_model(target_suffix: str) -> str:
    """
    Convert Target Model Suffix to Development Model
    Based on Excel logic
    """
    if not target_suffix:
        return ""

    left_12 = target_suffix[:12]
    mid_14_2 = target_suffix[13:15]   # MID(14,2)
    mid_13_2 = target_suffix[12:14]   # MID(13,2)
    mid_13_3 = target_suffix[12:15]   # MID(13,3)
    mid_16_1 = target_suffix[15:16]   # MID(16,1)
    mid_17_1 = target_suffix[16:17]   # MID(17,1)
    right_2 = target_suffix[-2:]

    # ==============================
    # EXCEL IF LOGIC (URUTAN PENTING)
    # ==============================

    # 1. TT
    if "TT" in mid_14_2:
        return f"{left_12}ATT{mid_16_1}LCX"

    # 2. JP / JE + 6
    if right_2 in ("JP", "JE") and mid_17_1 == "6":
        return f"{left_12}AEU{mid_16_1}6CX"

    # 3. JP / JE
    if right_2 in ("JP", "JE"):
        return f"{left_12}AEU{mid_16_1}LCX"

    # 4. WE / JE / JL + 6
    if right_2 in ("WE", "JE", "JL") and mid_17_1 == "6":
        return f"{left_12}AGM{mid_16_1}6CX"

    # 5. WE / JE / JL
    if right_2 in ("WE", "JE", "JL"):
        return f"{left_12}AGM{mid_16_1}LCX"

    # 6. SD / KD / SR / KR + 6
    if mid_13_2 in ("SD", "KD", "SR", "KR") and mid_17_1 == "6":
        return f"{left_12}AEU{mid_16_1}6CX"

    # 7. SD / KD / SR / KR / KNR
    if mid_13_2 in ("SD", "KD", "SR", "KR") or mid_13_3 == "KNR":
        return f"{left_12}AEU{mid_16_1}LCX"

    # 8. KT / ST / KAU / KPH / .NSB
    if any(x in target_suffix for x in ("KT", "ST", "KAU", "KPH", ".NSB")):
        return f"{left_12}AWZ{mid_16_1}LCX"

    # 9. KW / WN
    if any(x in target_suffix for x in ("KW", "WN")):
        return f"{left_12}AWZ{mid_16_1}LCX"

    # 10. HK
    if "HK" in target_suffix:
        return f"{left_12}AHK{mid_16_1}LCX"

    # 11. SW
    if "SW" in target_suffix:
        return f"{left_12}AWZ{mid_16_1}LCX"

    # 12. MW
    if "MW" in target_suffix:
        return f"{left_12}AWZ{mid_16_1}LCX"

    # 13. US
    if "US" in target_suffix:
        return f"{left_12}AUS{mid_16_1}LCX"

    # 14. CC
    if "CC" in target_suffix:
        return f"{left_12}AUS{mid_16_1}LCX"

    # 15. JL
    if "JL" in target_suffix:
        return f"{left_12}AJL{mid_16_1}LCX"

    # 16. KR
    if "KR" in target_suffix:
        return f"{left_12}AKR{mid_16_1}LCX"

    return ""


def normalize_model_input(value: str) -> str:
    """
    Normalize user input.
    Converts to uppercase and removes leading/trailing spaces.
    """
    if value is None:
        return ""
    return str(value).strip().upper()


def is_valid_model_input(value: str) -> bool:
    """
    Validate model text input.
    """
    if value is None:
        return False

    value = str(value).strip()

    if not value:
        return False

    invalid_values = {
        "Model Not Found",
        "Model Not Found (insert manually)",
        "Prediction Failed",
        "Prediction Failed (insert manually)",
        "Invalid Model",
        "Invalid Model (insert manually)",
    }

    if value in invalid_values:
        return False

    if len(value) < 8:
        return False

    if " " in value:
        return False

    return True


def get_initial(deve_model: str) -> str:
    if deve_model is None or str(deve_model).strip() == "":
        return ""

    InitialURL = (
        "https://api-nims-dev.apps.hedej.lge.com/api/custom/bom_crosscheck"
        f"?model_suffix={deve_model}"
    )

    try:
        response = requests.get(InitialURL, timeout=20)
        response.raise_for_status()

        getData = response.json()

        if not isinstance(getData, list) or len(getData) == 0:
            return ""

        parent_model_suffix = getData[0].get("parent_model_suffix", "")

        return normalize_model_input(parent_model_suffix)

    except Exception:
        return ""


error_list = []


def safe_get_initial(deve_model):
    if deve_model is None or str(deve_model).strip() == "":
        return ""

    try:
        return get_initial(str(deve_model).strip())

    except Exception as e:
        error_list.append({
            "dev model": deve_model,
            "error": str(e)
        })
        return ""


def predict_development(target_model: str) -> dict:
    """
    Predicts development model without streamlit dependency.
    Returns a dictionary containing prediction results and status messages.
    """
    norm_target = normalize_model_input(target_model)

    if not is_valid_model_input(norm_target):
        return {
            "dev_model": "",
            "manual_required": True,
            "message": "⚠️ Please input a valid Target Model first."
        }

    predicted_dev = normalize_model_input(get_development_model(norm_target))

    if is_valid_model_input(predicted_dev):
        return {
            "dev_model": predicted_dev,
            "manual_required": False,
            "message": ""
        }
    else:
        return {
            "dev_model": "",
            "manual_required": True,
            "message": "⚠️ Development Model could not be predicted. Please input it manually."
        }


def predict_initial(dev_model: str) -> dict:
    """
    Predicts initial model without streamlit dependency.
    Returns a dictionary containing prediction results and status messages.
    """
    norm_dev = normalize_model_input(dev_model)

    if not is_valid_model_input(norm_dev):
        return {
            "init_model": "",
            "manual_required": True,
            "message": "⚠️ Please input a valid Development Model before predicting Initial Model."
        }

    try:
        initial_model = normalize_model_input(get_initial(norm_dev))

        if is_valid_model_input(initial_model):
            return {
                "init_model": initial_model,
                "manual_required": False,
                "message": ""
            }
        else:
            return {
                "init_model": "",
                "manual_required": True,
                "message": "⚠️ Initial Model could not be found. Please input it manually."
            }

    except Exception as e:
        return {
            "init_model": "",
            "manual_required": True,
            "message": f"Failed to predict initial model. Please input Initial Model manually. Error: {e}"
        }


def predict_models(target_model: str) -> dict:
    """
    Orchestrates both development and initial model predictions.
    """
    dev_result = predict_development(target_model)
    
    if dev_result["manual_required"] or not dev_result["dev_model"]:
        return {
            "dev_model": dev_result["dev_model"],
            "dev_manual_required": dev_result["manual_required"],
            "dev_message": dev_result["message"],
            "init_model": "",
            "init_manual_required": True,
            "init_message": "⚠️ Skipped Initial Model prediction due to invalid/missing Development Model."
        }

    init_result = predict_initial_model(dev_result["dev_model"])

    return {
        "dev_model": dev_result["dev_model"],
        "dev_manual_required": dev_result["manual_required"],
        "dev_message": dev_result["message"],
        "init_model": init_result["init_model"],
        "init_manual_required": init_result["manual_required"],
        "init_message": init_result["message"]
    }

def predict_models_callback():
    predict_development()
    predict_initial_model()