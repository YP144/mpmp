
import io
import json
import re
import zipfile
from copy import copy as copy_style
from dataclasses import dataclass
from concurrent.futures import ThreadPoolExecutor, as_completed

import numpy as np
import pandas as pd
import polars as pl
import requests

from openpyxl import Workbook, load_workbook


from services.check_sso import SSO_NPDM
from services.main_hwpart import targetpart_compare, developmentpart_compare, initialpart_compare
from services.main_tooloption import run_tooloption, getToolinfo
from services.subs_hwpart import run_subshwpart

from functions.profile import ModelProfile, GetOperationPart
from functions.substate import GoGetTemplate, GetParseStructure, StructureCheck
from functions.allcompare import create_download_all_excel
from functions.datasend import upload_bytes_tos3
from functions.conditioning import zerobom_conditioning
from functions.tool_option_logic import GetModelToolOptionMaster, ToolOptionCross


from utils.dataframe_hash import get_dataframe_hash, get_dataframe_hashAgGrid, get_dataframe_hashStyler


def safe_load_workbook(excel_bytes):
    """Try loading as-is; if openpyxl chokes on conditional formatting, retry
    with conditional formatting stripped. Returns None if both attempts fail,
    so the caller can skip that source instead of crashing."""
    try:
        return load_workbook(io.BytesIO(excel_bytes))
    except (IndexError, KeyError):
        try:
            cleaned = strip_conditional_formatting(excel_bytes)
            return load_workbook(io.BytesIO(cleaned))
        except Exception:
            return None


def build_combined_workbook(sources):
    available = [(b, label) for b, label in sources if b]
    if not available:
        return None

    master_wb = Workbook()
    master_wb.remove(master_wb.active)

    skipped = []

    for excel_bytes, label in available:
        source_wb = safe_load_workbook(excel_bytes)

        if source_wb is None:
            skipped.append(label)
            continue

        for sheet_name in source_wb.sheetnames:
            copy_excel_sheet(source_wb[sheet_name], master_wb, sheet_name or label)

    if not master_wb.sheetnames:
        return None

    if skipped:
        st.warning(f"Could not include in Download All: {', '.join(skipped)}")

    master_buffer = io.BytesIO()
    master_wb.save(master_buffer)
    master_buffer.seek(0)
    return master_buffer.getvalue()