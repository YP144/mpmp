
# PROPRIETARY AND CONFIDENTIAL
#
# This file and its contents are the proprietary property of LGERC HW ATSC 60Hz.
# Unauthorized copying, redistribution, modification, disclosure, or claim of
# ownership is strictly prohibited.
#
# Use of LLMs, AI agents, AI chatbots, or automated code tools to alter,
# redefine, disguise, or claim this source code as another party’s own work is
# prohibited and may result in legal action.
#
# Removal or deletion of this notice, whether intentional or unintentional,
# does not waive any rights and shall be considered a breach of the applicable
# repository license notice.
#
# ALL RIGHTS RESERVED.



#from streamlit_navigation_bar import st_navbar
import pandas as pd
import time
import requests
import json
import io
import numpy as np
from openpyxl import load_workbook, workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import polars as pl
from pathlib import Path
import os
from dej_sdk import bucket
import tempfile
import re
import polars as pl
import hashlib
from pandas.io.formats.style import Styler
import connectorx

from main_function.services.check_sso import SSO_NPDM
from main_function.services.main_hwpart import targetpart_compare, developmentpart_compare, initialpart_compare
from main_function.services.main_tooloption import run_tooloption, getToolinfo
from main_function.services.subs_hwpart import run_subshwpart

def create_download_all_excel(
    target_model,
    dev_model,
    init_model,
    ProfileFrame,
    SummaryFrame=None
):
    """
    Create one Excel file containing combined results from:
    sheet1 Summary/Profile,
    sheet2 Structure,
    sheet3 BEN Compare,
    sheet4 Substitute,
    sheet5 HW Tool Option,
    sheet6 Tool Option Check,
    sheet7 Operation Plan.
    """

    output = io.BytesIO()

    with pd.ExcelWriter(output, engine="xlsxwriter") as writer:
        workbook = writer.book

        # =========================
        # Common Excel Formats
        # =========================
        title_format = workbook.add_format({
            "bold": True,
            "font_size": 20,
            "align": "left",
            "valign": "vcenter"
        })

        section_title_format = workbook.add_format({
            "bold": True,
            "font_size": 12,
            "bg_color": "#B4C6E7",
            "font_color": "#000000",
            "align": "left",
            "valign": "vcenter",
            "border": 1
        })

        header_format = workbook.add_format({
            "bold": True,
            "bg_color": "#D9E1F2",
            "font_color": "#000000",
            "align": "center",
            "valign": "vcenter",
            "border": 1
        })

        normal_format = workbook.add_format({
            "align": "left",
            "valign": "vcenter",
            "border": 1
        })

        remarks_format = workbook.add_format({
            "align": "left",
            "valign": "top",
            "text_wrap": True,
            "border": 1
        })

        ok_format = workbook.add_format({
            "bg_color": "#00B050",
            "font_color": "#FFFFFF",
            "bold": True,
            "align": "center",
            "valign": "vcenter",
            "border": 1
        })

        ng_format = workbook.add_format({
            "bg_color": "#FF1300",
            "font_color": "#FFFFFF",
            "bold": True,
            "align": "center",
            "valign": "vcenter",
            "border": 1
        })

        check_format = workbook.add_format({
            "bg_color": "#F5C527",
            "font_color": "#000000",
            "bold": True,
            "align": "center",
            "valign": "vcenter",
            "border": 1
        })

        red_diff_format = workbook.add_format({
            "bg_color": "#FF1300",
            "font_color": "#FFFFFF",
            "bold": True,
            "align": "left",
            "valign": "vcenter",
            "border": 1
        })

        def safe_excel_value(value):
            if value is None:
                return ""

            try:
                if pd.isna(value):
                    return ""
            except Exception:
                pass

            return value

        def normalize_compare_value(value):
            if value is None:
                return ""

            try:
                if pd.isna(value):
                    return ""
            except Exception:
                pass

            value = str(value).strip()

            if value.lower() == "nan":
                return ""

            return value

        def to_pandas_safe(df):
            if df is None:
                return pd.DataFrame()

            if isinstance(df, pd.io.formats.style.Styler):
                df = df.data

            if isinstance(df, pl.DataFrame):
                df = df.to_pandas()

            if isinstance(df, pd.DataFrame):
                return df.copy()

            try:
                return pd.DataFrame(df)
            except Exception:
                return pd.DataFrame()

        def write_basic_table(
            worksheet,
            df,
            startrow,
            startcol,
            result_col_name=None,
            remarks_col_names=None
        ):
            df = to_pandas_safe(df)

            if df.empty:
                return startrow

            df.columns = [str(col).strip() for col in df.columns]
            columns = list(df.columns)

            if remarks_col_names is None:
                remarks_col_names = ["Remarks", "Al. Remarks", "Eng. Remarks"]

            for col_idx, col_name in enumerate(columns):
                worksheet.write(startrow, startcol + col_idx, col_name, header_format)

            for row_idx in range(len(df)):
                excel_row = startrow + 1 + row_idx

                for col_idx, col_name in enumerate(columns):
                    value = safe_excel_value(df.iloc[row_idx, col_idx])
                    cell_format = normal_format

                    if col_name in remarks_col_names:
                        cell_format = remarks_format

                    if result_col_name is not None and col_name == result_col_name:
                        result_value = str(value).strip().upper()

                        if result_value == "OK":
                            cell_format = ok_format
                        elif result_value == "NG":
                            cell_format = ng_format
                        elif result_value == "CHECK":
                            cell_format = check_format

                    worksheet.write(excel_row, startcol + col_idx, value, cell_format)

            for col_idx, col_name in enumerate(columns):
                max_len = len(str(col_name))

                if len(df) > 0:
                    max_len = max(max_len, df[col_name].astype(str).map(len).max())

                worksheet.set_column(
                    startcol + col_idx,
                    startcol + col_idx,
                    min(max_len + 3, 45)
                )

            return startrow + len(df) + 1

        def write_compare_table(
            worksheet,
            df,
            startrow,
            startcol,
            initial_col_name,
            compare_col_names,
            remarks_col_names=None
        ):
            df = to_pandas_safe(df)

            if df.empty:
                return startrow

            df.columns = [str(col).strip() for col in df.columns]
            columns = list(df.columns)

            if remarks_col_names is None:
                remarks_col_names = ["Remarks", "Al. Remarks", "Eng. Remarks"]

            for col_idx, col_name in enumerate(columns):
                worksheet.write(startrow, startcol + col_idx, col_name, header_format)

            initial_idx = columns.index(initial_col_name) if initial_col_name in columns else None

            for row_idx in range(len(df)):
                excel_row = startrow + 1 + row_idx

                initial_value = ""
                if initial_idx is not None:
                    initial_value = normalize_compare_value(df.iloc[row_idx, initial_idx])

                for col_idx, col_name in enumerate(columns):
                    raw_value = df.iloc[row_idx, col_idx]
                    cell_value = safe_excel_value(raw_value)
                    cell_format = normal_format

                    if col_name in remarks_col_names:
                        cell_format = remarks_format

                    if col_name in compare_col_names:
                        current_value = normalize_compare_value(raw_value)

                        if initial_value == "" or current_value == "" or current_value != initial_value:
                            cell_format = red_diff_format

                    worksheet.write(excel_row, startcol + col_idx, cell_value, cell_format)

            for col_idx, col_name in enumerate(columns):
                max_len = len(str(col_name))

                if len(df) > 0:
                    max_len = max(max_len, df[col_name].astype(str).map(len).max())

                worksheet.set_column(
                    startcol + col_idx,
                    startcol + col_idx,
                    min(max_len + 3, 45)
                )

            return startrow + len(df) + 1

        # =========================================================
        # Sheet 1: Summary / Model Profile
        # =========================================================
        worksheet = workbook.add_worksheet("Summary")
        writer.sheets["Summary"] = worksheet

        worksheet.merge_range(
            0,
            0,
            0,
            8,
            "Model Profile Summary - " + str(target_model),
            title_format
        )

        worksheet.write(2, 1, "Target Model", section_title_format)
        worksheet.write(2, 2, str(target_model), normal_format)

        worksheet.write(3, 1, "Development Model", section_title_format)
        worksheet.write(3, 2, str(dev_model), normal_format)

        worksheet.write(4, 1, "Initial Model", section_title_format)
        worksheet.write(4, 2, str(init_model), normal_format)

        worksheet.write(6, 1, "Model Profile", section_title_format)
        write_basic_table(
            worksheet=worksheet,
            df=ProfileFrame,
            startrow=7,
            startcol=1
        )

        if SummaryFrame is not None:
            worksheet.write(21, 1, "Summary Result", section_title_format)
            write_basic_table(
                worksheet=worksheet,
                df=SummaryFrame,
                startrow=22,
                startcol=1,
                result_col_name="Result"
            )

        if "tool_frame" in st.session_state:
            worksheet.write(32, 1, "Tool Option Summary", section_title_format)
            write_basic_table(
                worksheet=worksheet,
                df=st.session_state.tool_frame,
                startrow=33,
                startcol=1,
                result_col_name="Result"
            )

        worksheet.freeze_panes(7, 1)
        worksheet.set_zoom(90)

        # =========================================================
        # Sheet 2: Structure
        # =========================================================
        worksheet = workbook.add_worksheet("Structure")
        writer.sheets["Structure"] = worksheet

        worksheet.merge_range(
            0,
            0,
            0,
            8,
            "Structure Checklist - " + str(target_model),
            title_format
        )

        if "structure_frame" in st.session_state:
            write_basic_table(
                worksheet=worksheet,
                df=st.session_state.structure_frame,
                startrow=3,
                startcol=1
            )

        if "remark_frame" in st.session_state:
            remark_startrow = 3
            structure_df = to_pandas_safe(st.session_state.structure_frame)
            if not structure_df.empty:
                remark_startrow = len(structure_df) + 7

            worksheet.write(remark_startrow - 1, 1, "Structure Remarks", section_title_format)
            write_basic_table(
                worksheet=worksheet,
                df=st.session_state.remark_frame,
                startrow=remark_startrow,
                startcol=1
            )

        worksheet.set_zoom(90)

        # =========================================================
        # Sheet 3: BEN Compare
        # =========================================================
        worksheet = workbook.add_worksheet("BEN Compare")
        writer.sheets["BEN Compare"] = worksheet

        worksheet.merge_range(
            0,
            0,
            0,
            10,
            "BEN Compare - " + str(target_model),
            title_format
        )

        worksheet.write(1, 1, "Target Model", section_title_format)
        worksheet.write(1, 2, str(target_model), normal_format)
        worksheet.write(1, 3, "Development Model", section_title_format)
        worksheet.write(1, 4, str(dev_model), normal_format)
        worksheet.write(1, 5, "Initial Model", section_title_format)
        worksheet.write(1, 6, str(init_model), normal_format)

        ben_df = st.session_state.get("BEN_CompareRemarks_edited")
        if ben_df is not None:
            worksheet.write(3, 1, "BOM Part Compare", section_title_format)
            write_compare_table(
                worksheet=worksheet,
                df=ben_df,
                startrow=4,
                startcol=1,
                initial_col_name="Part No_Initial",
                compare_col_names=["Part No_Development", "Part No_Target"]
            )

        toolopt_df = st.session_state.get("ToolOptionRemarks_edited")
        if toolopt_df is not None:
            ben_df_len = len(to_pandas_safe(ben_df))
            tool_startrow = max(31, ben_df_len + 8)

            worksheet.write(tool_startrow - 1, 1, "Tool Option Compare", section_title_format)
            write_compare_table(
                worksheet=worksheet,
                df=toolopt_df,
                startrow=tool_startrow,
                startcol=1,
                initial_col_name="Initial",
                compare_col_names=["Development", "Target"]
            )

        worksheet.freeze_panes(4, 1)
        worksheet.set_zoom(90)

        # =========================================================
        # Sheet 4: Substitute & Part Status
        # =========================================================
        worksheet = workbook.add_worksheet("Substitute")
        writer.sheets["Substitute"] = worksheet

        worksheet.merge_range(
            0,
            0,
            0,
            10,
            "Substitute & Part Status - " + str(target_model),
            title_format
        )

        status_df = st.session_state.get("compare_result_statuspart")
        subs_df = st.session_state.get("compare_result_substitute")

        if status_df is not None:
            status_df = to_pandas_safe(status_df)

            if "Al. Remarks" not in status_df.columns:
                status_df["Al. Remarks"] = ""

            if "Eng. Remarks" not in status_df.columns:
                status_df["Eng. Remarks"] = ""

            worksheet.write(3, 1, "Part Status Compare", section_title_format)
            write_compare_table(
                worksheet=worksheet,
                df=status_df,
                startrow=4,
                startcol=1,
                initial_col_name="Status_Init",
                compare_col_names=["Status_Dev", "Status_Target"]
            )

        if subs_df is not None:
            subs_df = to_pandas_safe(subs_df)

            if "Al. Remarks" not in subs_df.columns:
                subs_df["Al. Remarks"] = ""

            if "Eng. Remarks" not in subs_df.columns:
                subs_df["Eng. Remarks"] = ""

            status_len = len(to_pandas_safe(status_df))
            subs_startrow = status_len + 9

            worksheet.write(subs_startrow - 1, 1, "Substitute Logic Compare", section_title_format)
            write_compare_table(
                worksheet=worksheet,
                df=subs_df,
                startrow=subs_startrow,
                startcol=1,
                initial_col_name="Subs_Init",
                compare_col_names=["Subs_Dev", "Subs_Target"]
            )

        worksheet.set_zoom(90)

        # =========================================================
        # Sheet 5: HW Tool Option Compare
        # =========================================================
        worksheet = workbook.add_worksheet("HW Tool Option")
        writer.sheets["HW Tool Option"] = worksheet

        worksheet.merge_range(
            0,
            0,
            0,
            10,
            "HW Tool Option Compare - " + str(target_model),
            title_format
        )

        single_df = st.session_state.get("HWSingleOptRemarks_edited")
        multi_df = st.session_state.get("HWMultiOptRemarks_edited")

        if single_df is not None:
            worksheet.write(3, 1, "Single Option Compare", section_title_format)
            write_compare_table(
                worksheet=worksheet,
                df=single_df,
                startrow=4,
                startcol=1,
                initial_col_name="Initial",
                compare_col_names=["Development", "Target"],
                remarks_col_names=["Eng. Remarks"]
            )

        if multi_df is not None:
            single_len = len(to_pandas_safe(single_df))
            multi_startrow = max(61, single_len + 9)

            worksheet.write(multi_startrow - 1, 1, "Multi Option Compare", section_title_format)
            write_compare_table(
                worksheet=worksheet,
                df=multi_df,
                startrow=multi_startrow,
                startcol=1,
                initial_col_name="Initial",
                compare_col_names=["Development", "Target"],
                remarks_col_names=["Eng. Remarks"]
            )

        worksheet.set_zoom(90)

        # =========================================================
        # Sheet 6: Tool Option Check
        # =========================================================
        worksheet = workbook.add_worksheet("Tool Option Check")
        writer.sheets["Tool Option Check"] = worksheet

        worksheet.merge_range(
            0,
            0,
            0,
            10,
            "Tool Option Check - " + str(target_model),
            title_format
        )

        if "tool_frame" in st.session_state:
            worksheet.write(3, 1, "Tool Option Summary", section_title_format)
            write_basic_table(
                worksheet=worksheet,
                df=st.session_state.tool_frame,
                startrow=4,
                startcol=1,
                result_col_name="Result"
            )

        if "tool_multi_frame" in st.session_state:
            tool_frame_len = len(to_pandas_safe(st.session_state.tool_frame))
            multi_tool_startrow = tool_frame_len + 9

            worksheet.write(multi_tool_startrow - 1, 1, "Tool Option Master Checklist", section_title_format)
            write_basic_table(
                worksheet=worksheet,
                df=st.session_state.tool_multi_frame,
                startrow=multi_tool_startrow,
                startcol=1,
                result_col_name="Result"
            )

        worksheet.set_zoom(90)

        # =========================================================
        # Sheet 7: Operation Plan
        # =========================================================
        worksheet = workbook.add_worksheet("Operation Plan")
        writer.sheets["Operation Plan"] = worksheet

        worksheet.merge_range(
            0,
            0,
            0,
            10,
            "Operation Plan Check - " + str(target_model),
            title_format
        )

        if "op_frame" in st.session_state:
            write_basic_table(
                worksheet=worksheet,
                df=st.session_state.op_frame,
                startrow=3,
                startcol=1,
                result_col_name="Result"
            )

        worksheet.set_zoom(90)

    output.seek(0)
    return output
